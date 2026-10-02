"""Drawing kit for the Quant-A-Maze 3.O deck: template cloning plus a small set of native PowerPoint primitives.

Every coordinate is in inches and converted to an integer EMU at the point of use: real PowerPoint refuses a file
with a float coordinate (`x="12641000.0"`) even though every lenient checker accepts it.
"""
from __future__ import annotations

import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ---------------------------------------------------------------- palette and type (from the official template)
ORANGE = "FF914D"      # the template's band and halftone orange: fills and lines only
DEEP = "C2410C"        # orange for text on white (contrast about 5:1)
TINT = "FFE9DB"        # light orange fill
INK = "000000"
TEXT = "1A1A1A"
MUTED = "5C5C5C"
RULE = "BDBDBD"
CARD = "F6F6F6"
DARK = "1F1F1F"
WHITE = "FFFFFF"

TITLE_FONT = "Bahnschrift SemiBold"
BODY_FONT = "Segoe UI"
STRONG_FONT = "Segoe UI Semibold"
MONO_FONT = "Consolas"

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Content frame inside the template chrome (inches): header band ends at 1.83, halftone strip starts at ~17.5,
# bottom band starts at 10.29.
LEFT, RIGHT, TOP, BOTTOM = 0.9, 17.1, 3.75, 10.0


def E(v: float) -> Emu:
    return Emu(int(round(v * 914400)))


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


# ---------------------------------------------------------------- template handling
def clone_slide(prs, src, keep=lambda shape: True):
    """Append a copy of `src` (shapes and image relationships) and return it. `keep(shape)` filters shapes."""
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag in (qn("p:sp"), qn("p:grpSp"), qn("p:pic"), qn("p:cxnSp"), qn("p:graphicFrame")):
            tree.remove(el)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rid_map[rid] = new.part.rels.get_or_add(rel.reltype, rel.target_part)
    for shape in src.shapes:
        if not keep(shape):
            continue
        el = copy.deepcopy(shape._element)
        for node in el.iter():
            for attr in list(node.attrib):
                if attr.startswith("{%s}" % R_NS) and node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        tree.append(el)
    return new


def drop_slides(prs, count: int) -> None:
    """Remove the first `count` slides (the template's own pages) from the presentation."""
    ids = prs.slides._sldIdLst
    for sld in list(ids)[:count]:
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


# ---------------------------------------------------------------- text
def _style_run(run, size, color, font, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def _bullet(paragraph, char="■", color=ORANGE, indent=0.32, level=0):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(int(round((indent + level * 0.35) * 914400))))
    ppr.set("indent", str(int(round(-indent * 914400))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)
    clr = etree.SubElement(ppr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(ppr, qn("a:buSzPct")).set("val", "70000")
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(ppr, qn("a:buChar")).set("char", char)


def text(slide, x, y, w, h, paras, size=22, color=TEXT, font=BODY_FONT, bold=False, align="l", anchor="t",
         spacing=1.08, after=4, bullets=False, inset=0.0):
    """paras: a string, or a list of paragraphs; a paragraph is a string or a list of (text, style dict) runs.
    A paragraph may also be a dict {"runs": [...], "bullet": bool, "level": int, "size": int, "after": pt}."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, E(inset))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    body = tf._txBody
    bp = body.find(qn("a:bodyPr"))
    for old in list(bp):
        bp.remove(old)
    etree.SubElement(bp, qn("a:noAutofit"))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        spec = para if isinstance(para, dict) else {"runs": para}
        runs = spec["runs"]
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[spec.get("align", align)]
        p.line_spacing = spacing
        p.space_after = Pt(spec.get("after", after))
        if spec.get("before"):
            p.space_before = Pt(spec["before"])
        if spec.get("bullet", bullets):
            _bullet(p, level=spec.get("level", 0), char=spec.get("char", "■"), color=spec.get("bullet_color", ORANGE))
        for txt, st in runs:
            r = p.add_run()
            r.text = txt
            _style_run(r, st.get("size", spec.get("size", size)), st.get("color", spec.get("color", color)),
                       st.get("font", spec.get("font", font)), st.get("bold", spec.get("bold", bold)),
                       st.get("italic", False))
    return tb


# ---------------------------------------------------------------- shapes
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
          "hex": MSO_SHAPE.HEXAGON, "diamond": MSO_SHAPE.DIAMOND, "chevron": MSO_SHAPE.CHEVRON,
          "pentagon": MSO_SHAPE.PENTAGON}


def box(slide, x, y, w, h, label=None, sub=None, fill=WHITE, line=INK, lw=1.5, kind="round", radius=0.12,
        size=20, sub_size=16, color=TEXT, sub_color=MUTED, font=STRONG_FONT, sub_font=BODY_FONT, align="c",
        anchor="m", dash=False, inset=0.1, shadow=False):
    shp = slide.shapes.add_shape(SHAPES[kind], E(x), E(y), E(w), E(h))
    if kind == "round":
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            ln = shp.line._get_or_add_ln()
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if not shadow:
        sp_pr = shp._element.spPr
        etree.SubElement(sp_pr, qn("a:effectLst"))
    tf = shp.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, E(inset))
    tf.margin_top = tf.margin_bottom = E(0.04)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    bp = tf._txBody.find(qn("a:bodyPr"))
    etree.SubElement(bp, qn("a:noAutofit"))
    first = True
    for content, sz, col, fnt in ((label, size, color, font), (sub, sub_size, sub_color, sub_font)):
        if content is None:
            continue
        for line_text in content.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = 1.0
            r = p.add_run()
            r.text = line_text
            _style_run(r, sz, col, fnt)
    return shp


def node(slide, cx, cy, r=0.09, fill=ORANGE, line=None, lw=1.25):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - r), E(cy - r), E(2 * r), E(2 * r))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    etree.SubElement(shp._element.spPr, qn("a:effectLst"))
    return shp


def line(slide, x1, y1, x2, y2, color=INK, lw=1.5, arrow=True, dash=False, start_arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if start_arrow:
        h = etree.SubElement(ln, qn("a:headEnd"))
        h.set("type", "triangle"); h.set("w", "med"); h.set("len", "med")
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


def elbow(slide, points, color=INK, lw=1.5, arrow=True, dash=False):
    """A polyline of straight segments; the arrow head only on the last one."""
    for i in range(len(points) - 1):
        (x1, y1), (x2, y2) = points[i], points[i + 1]
        line(slide, x1, y1, x2, y2, color=color, lw=lw, arrow=arrow and i == len(points) - 2, dash=dash)


def rule(slide, x, y, w, color=RULE, lw=1.0):
    return line(slide, x, y, x + w, y, color=color, lw=lw, arrow=False)


def chip(slide, x, y, label, w=None, kind="real", size=15):
    """Small status tag. kind: real (black outline), sim (orange outline), next (grey dashed), hot (orange fill)."""
    w = w or (0.2 + 0.105 * len(label) * size / 15)
    style = {"real": (WHITE, INK, INK, False), "sim": (WHITE, ORANGE, DEEP, False), "next": (WHITE, MUTED, MUTED, True),
             "hot": (ORANGE, ORANGE, INK, False), "dark": (DARK, DARK, WHITE, False)}[kind]
    return box(slide, x, y, w, 0.38, label, fill=style[0], line=style[1], color=style[2], size=size, lw=1.25,
               radius=0.5, dash=style[3], inset=0.06)


def notes(slide, text_):
    slide.notes_slide.notes_text_frame.text = text_


# ---------------------------------------------------------------- slide furniture
def header(slide, section, title, subtitle=None, official=False):
    """section: the official template heading this slide belongs to. On a section-opening slide the title IS the
    official heading (official=True) and the section marker is not repeated."""
    if not official:
        text(slide, LEFT, 1.98, 12.0, 0.36, [[(section, {"color": DEEP})]], size=16, font=STRONG_FONT)
    text(slide, LEFT, 2.28 if not official else 2.05, 16.2, 0.8, title, size=40, color=INK, font=TITLE_FONT)
    top = 3.02 if not official else 2.82
    if subtitle:
        text(slide, LEFT, top, 16.2, 0.5, subtitle, size=21, color=MUTED)
    # the template's "o——" motif under the title block
    yy = top + (0.6 if subtitle else 0.12)
    node(slide, LEFT + 0.07, yy, r=0.07, fill=WHITE, line=ORANGE, lw=1.75)
    line(slide, LEFT + 0.14, yy, LEFT + 1.6, yy, color=ORANGE, lw=1.75, arrow=False)


def footer(slide, number, total):
    text(slide, LEFT, 10.66, 9.0, 0.5, [[("Q-SHIELD", {"font": STRONG_FONT}),
                                        ("   continuous device trust with post-quantum-signed evidence", {})]],
         size=15, color=INK, anchor="m")
    text(slide, 14.6, 10.66, 2.5, 0.5, f"{number} / {total}", size=15, color=INK, align="r", anchor="m")
