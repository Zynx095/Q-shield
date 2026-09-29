"""Drawing kit for the final deck: native (editable) PowerPoint shapes with one consistent visual language.

Status language (shape + colour + text, never colour alone):
    IMPLEMENTED  green dot      DESIGNED  violet diamond      NEXT  amber chevron
Two font families only: Segoe UI (text) and Consolas (technical labels, numbers).
"""
import math
from dataclasses import dataclass

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SANS, MONO = "Segoe UI", "Consolas"
BG, PANEL, PANEL2 = "05070D", "0D1424", "111B30"
LINE, LINE2 = "1F2C4A", "33466F"
TXT, MUT, DIM = "F1F5F9", "9AA9C0", "66768F"
CY, VI, BL, AM, GN, RD = "22D3EE", "8B5CF6", "7DB4FF", "FBBF24", "34D399", "F87171"
STATUS = {"IMPL": ("IMPLEMENTED", GN, MSO_SHAPE.OVAL), "DES": ("DESIGNED", VI, MSO_SHAPE.DIAMOND),
          "NEXT": ("NEXT", AM, MSO_SHAPE.CHEVRON)}


def rgb(h):
    return RGBColor.from_string(h)


@dataclass
class N:
    """A placed node: centre + size, used to route connectors to rectangle edges."""
    cx: float
    cy: float
    w: float
    h: float

    @property
    def x(self):
        return self.cx - self.w / 2

    @property
    def y(self):
        return self.cy - self.h / 2


def _alpha(fill_owner, pct):
    if pct < 100:
        clr = fill_owner._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(pct * 1000)))


def _effect_list(shape_element):
    """The shape's single a:effectLst (python-pptx may already have created an empty one)."""
    sp = shape_element.spPr
    eff = sp.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(sp, qn("a:effectLst"))
    return eff


def glow(shape, hexc, rad=9, alpha=40):
    eff = _effect_list(shape._element)
    g = etree.SubElement(eff, qn("a:glow"))
    g.set("rad", str(int(rad * 12700)))
    c = etree.SubElement(g, qn("a:srgbClr"))
    c.set("val", hexc)
    etree.SubElement(c, qn("a:alpha")).set("val", str(alpha * 1000))


def _runs(p, runs, size, bold, color, font, italic=False, spc=0):
    if isinstance(runs, str):
        runs = [(runs, {})]
    for t, o in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic)
        f.color.rgb = rgb(o.get("color", color))
        s = o.get("spc", spc)
        if s:
            r._r.get_or_add_rPr().set("spc", str(int(s)))


def fill_text(tf, paras, size=12, bold=False, color=TXT, font=SANS, align="l", anchor="t", margins=(0, 0, 0, 0),
              spc=0, italic=False, space_after=0, line_spacing=None):
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str) or (isinstance(paras, list) and paras and isinstance(paras[0], tuple)):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        _runs(p, para, size, bold, color, font, italic, spc)


def text(s, x, y, w, h, paras, **kw):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(box.text_frame, paras, **kw)
    return box


def shape(s, kind, x, y, w, h, fill=PANEL, alpha=100, line=None, lw=1.0, dashed=False, rad=None, glow_=None):
    shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE and rad is not None:
        shp.adjustments[0] = min(0.5, rad / min(w, h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        _alpha(shp.fill, alpha)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dashed:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if glow_:
        glow(shp, glow_[0], glow_[1], glow_[2])
    return shp


def panel(s, x, y, w, h, line=LINE2, fill=PANEL, alpha=70, rad=0.14, dashed=False, lw=1.0, glow_=None):
    return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, alpha=alpha, line=line, lw=lw, dashed=dashed, rad=rad, glow_=glow_)


def node(s, x, y, w, h, title, sub=None, accent=CY, fill=PANEL2, alpha=95, tsize=12, ssize=10, dashed=False, glow_=None,
         lw=1.25, tcolor=TXT, kind=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.12, status=None, mono=False):
    """Node with centred title (+ optional sub line and status line). Returns N for routing."""
    shp = shape(s, kind, x, y, w, h, fill=fill, alpha=alpha, line=accent, lw=lw, dashed=dashed, rad=rad, glow_=glow_)
    paras = [[(title, {"size": tsize, "bold": True, "color": tcolor, "font": MONO if mono else SANS})]]
    if sub:
        paras.append([(sub, {"size": ssize, "color": MUT})])
    if status:
        label, col, _ = STATUS[status]
        paras.append([(label, {"size": 10, "color": col, "font": MONO, "bold": True})])
    fill_text(shp.text_frame, paras, align="c", anchor="m", margins=(0.05, 0.02, 0.05, 0.02))
    return N(x + w / 2, y + h / 2, w, h)


def edge_point(n: N, tx, ty, pad=0.04):
    dx, dy = tx - n.cx, ty - n.cy
    if dx == 0 and dy == 0:
        return n.cx, n.cy
    t = min((n.w / 2) / abs(dx) if dx else math.inf, (n.h / 2) / abs(dy) if dy else math.inf)
    L = math.hypot(dx, dy)
    ex, ey = n.cx + dx * t, n.cy + dy * t
    return ex + dx / L * pad, ey + dy / L * pad


def link(s, p1, p2, color=CY, w=1.5, dashed=False, arrow="end", glow_=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]), Inches(p2[0]), Inches(p2[1]))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    if arrow in ("start", "both"):
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle"), e.set("w", "med"), e.set("len", "med")
    if arrow in ("end", "both"):
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle"), e.set("w", "med"), e.set("len", "med")
    if glow_:
        eff = _effect_list(c._element)
        g = etree.SubElement(eff, qn("a:glow"))
        g.set("rad", str(8 * 12700))
        clr = etree.SubElement(g, qn("a:srgbClr"))
        clr.set("val", color)
        etree.SubElement(clr, qn("a:alpha")).set("val", "35000")
    return c


def connect(s, a: N, b: N, **kw):
    p1, p2 = edge_point(a, b.cx, b.cy), edge_point(b, a.cx, a.cy)
    return link(s, p1, p2, **kw)


def polyline(s, pts, color=CY, w=1.5, dashed=False, arrow=True):
    for i in range(len(pts) - 1):
        link(s, pts[i], pts[i + 1], color=color, w=w, dashed=dashed, arrow="end" if (arrow and i == len(pts) - 2) else None)


def dots(s, p1, p2, n=3, color=CY, size=0.075):
    for i in range(1, n + 1):
        t = i / (n + 1)
        x, y = p1[0] + (p2[0] - p1[0]) * t, p1[1] + (p2[1] - p1[1]) * t
        shape(s, MSO_SHAPE.OVAL, x - size / 2, y - size / 2, size, size, fill=color, glow_=(color, 5, 45))


def tag(s, x, y, kind, align="l", size=10):
    """Status mark (shape) + status text label. Returns approximate width."""
    label, col, mark = STATUS[kind]
    wtxt = 0.085 * len(label) + 0.05
    m = 0.13
    if align == "l":
        shape(s, mark, x, y + 0.035, m, m, fill=col)
        text(s, x + m + 0.07, y, wtxt, 0.2, [(label, {"color": col, "bold": True})], size=size, font=MONO, anchor="m")
    return m + 0.07 + wtxt


def chip(s, x, y, w, h, label, color=CY, fill=PANEL2, size=10.5, mono=True, bold=True, alpha=90, tcolor=None):
    shp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, alpha=alpha, line=color, lw=1.0, rad=h / 2)
    fill_text(shp.text_frame, [(label, {"color": tcolor or color, "bold": bold})], size=size, font=MONO if mono else SANS,
              align="c", anchor="m", margins=(0.04, 0, 0.04, 0))
    return N(x + w / 2, y + h / 2, w, h)


def legend(s, x, y, gap=0.25):
    cx = x
    for k in ("IMPL", "DES", "NEXT"):
        cx += tag(s, cx, y, k) + gap
    return cx
