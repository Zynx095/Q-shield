"""Build the ORIGINAL 6-slide Q-SHIELD idea submission deck (designed from scratch; not derived from the sample).

    python ppt/tools/final_deck/build_final_deck.py [--team-name ... --members ... --school ... --contact ...
                                                      --ps-number ... --track ... --title ...] [--out ...]

The official sample is used only for the constraints: max 6 slides, points and diagrams only, and the required
section headings (kept verbatim): IDEA TITLE / TECHNICAL APPROACH / FEASIBILITY AND VIABILITY /
IMPACT AND BENEFITS / RESEARCH AND REFERENCES, plus the title-slide fields. Every diagram is built from native,
editable PowerPoint shapes. Facts come only from ppt/references/verified-facts.md.
"""
import argparse
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
import backgrounds  # noqa: E402
from kit import *  # noqa: E402,F401,F403

ROOT = Path(__file__).resolve().parents[3]
BGDIR = ROOT / "ppt" / "final" / "backgrounds"
TOTAL = 6


def new_slide(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(BGDIR / f"bg{n}.jpg"), 0, 0, prs.slide_width, prs.slide_height)
    return s


def chrome(s, n, heading, team, extra=None):
    text(s, 0.6, 0.28, 12.1, 0.32, [(f"{n:02d}", {"color": CY, "font": MONO, "bold": True, "size": 12}),
                                    ("    " + heading, {"bold": True, "size": 13, "spc": 120}),
                                    ] + ([("      " + extra, {"color": MUT, "size": 12})] if extra else []), anchor="m")
    text(s, 0.6, 7.1, 8, 0.24, [("Q-SHIELD", {"bold": True, "color": TXT}), (f"   ·   TEAM  {team}", {"color": DIM})], size=10, font=MONO, anchor="m")
    text(s, 11.2, 7.1, 1.53, 0.24, [(f"{n:02d} / {TOTAL:02d}", {"color": DIM})], size=10, font=MONO, align="r", anchor="m")


def ring(cx, cy, rx, ry, k, n, start=-90):
    a = math.radians(start + k * 360 / n)
    return cx + rx * math.cos(a), cy + ry * math.sin(a)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------------------------------------ slide 1
def slide1(prs, a):
    s = new_slide(prs, 1)
    text(s, 0.6, 0.36, 12.1, 0.26, [("MINI HACKATHON  ·  23 SEPTEMBER 2026  ·  NVIDIA AI LAB  ·  THE DIGITAL MISSION 2026  —  IDEA SUBMISSION", {"color": DIM})],
         size=10, font=MONO, spc=60)
    text(s, 0.6, 0.85, 6.4, 0.3, [("CYBERSECURITY  ×  AI  ×  POST-QUANTUM  ×  AIoT", {"color": CY, "bold": True})], size=12, font=MONO, spc=100)
    text(s, 0.6, 1.2, 6.4, 1.5, [[("Q-", {"color": CY}), ("SHIELD", {"color": TXT})]], size=84, bold=True, anchor="m")
    text(s, 0.6, 2.75, 6.2, 0.95, ["Quantum-Resilient Self-Healing", "AIoT Security Architecture"], size=23, bold=True, color=TXT, line_spacing=0.95)
    text(s, 0.6, 3.7, 6.0, 0.6, [("Continuous device trust for AIoT systems in a post-quantum world.", {"color": MUT})], size=14)
    legend(s, 0.6, 4.4)
    # required title-slide fields, as a compact brief panel
    panel(s, 0.6, 4.85, 6.05, 2.05, alpha=55, line=LINE)
    cells = [("PROBLEM STATEMENT NUMBER", a.ps_number, 0.8, 4.95, 1.9), ("PROBLEM STATEMENT TITLE", a.title, 3.05, 4.95, 3.5),
             ("TRACK / THEME", a.track, 0.8, 5.53, 1.9), ("TEAM NAME", a.team_name, 2.75, 5.53, 1.9), ("SCHOOL", a.school, 4.7, 5.53, 1.85),
             ("TEAM MEMBERS (4-MAX)", a.members, 0.8, 6.11, 3.85), ("TEAM LEADER CONTACT", a.contact, 4.7, 6.11, 1.85)]
    for lab, val, x, y, w in cells:
        text(s, x, y, w, 0.2, [(lab, {"color": DIM})], size=10, font=MONO)
        text(s, x, y + 0.2, w, 0.3, [(val, {"color": TXT, "bold": True})], size=12)
    # hero visual: the Q-SHIELD core and the signals that feed it
    cx, cy = 10.0, 3.7
    for r, d in ((1.15, False), (2.0, True), (2.85, True)):
        shape(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill=None, line=LINE2, lw=0.75, dashed=d)
    for ang in range(0, 360, 40):                                     # faint topology nodes on the outer ring
        px, py = cx + 2.85 * math.cos(math.radians(ang + 15)), cy + 2.85 * math.sin(math.radians(ang + 15))
        shape(s, MSO_SHAPE.OVAL, px - 0.05, py - 0.05, 0.1, 0.1, fill=LINE2)
    sats = {"cam": node(s, 7.55, 1.05, 1.9, 0.72, "CAMERA", "USB · YOLO11n", accent=GN, tsize=12),
            "esp": node(s, 6.75, 3.35, 1.9, 0.72, "ESP32", "HMAC-SHA256", accent=VI, tsize=12),
            "pqc": node(s, 11.0, 2.05, 1.7, 0.72, "PQC", "ML-KEM · ML-DSA", accent=GN, tsize=12),
            "ai": node(s, 10.55, 5.55, 2.15, 0.72, "AI + TRUST", "dynamic trust", accent=AM, tsize=12, dashed=True)}
    core = N(cx, cy, 1.5, 1.3)
    for k, sn in sats.items():
        p1, p2 = edge_point(sn, cx, cy, 0.03), edge_point(core, sn.cx, sn.cy, 0.03)
        link(s, p1, p2, color=AM if k == "ai" else CY, w=1.6, dashed=(k == "ai"), arrow="both", glow_=(k != "ai"))
        if k != "ai":
            dots(s, p1, p2, 2)
    hexa = shape(s, MSO_SHAPE.HEXAGON, cx - 0.85, cy - 0.75, 1.7, 1.5, fill="0A2433", line=CY, lw=2.25, glow_=(CY, 16, 45))
    fill_text(hexa.text_frame, [[("Q-SHIELD", {"size": 14, "bold": True})], [("CORE", {"size": 10, "color": CY, "font": MONO, "bold": True})]],
              align="c", anchor="m")
    shape(s, MSO_SHAPE.HEXAGON, cx - 0.6, cy - 0.53, 1.2, 1.06, fill=None, line=CY, lw=0.6)
    for k, lab, (x, y) in (("cam", "signed observations", (6.85, 2.3)), ("esp", "device telemetry", (6.85, 4.2)), ("pqc", "session keys", (11.15, 3.15))):
        text(s, x, y, 1.9, 0.22, [(lab, {"color": DIM})], size=10, font=MONO)
    notes(s, "Q-SHIELD: continuous device trust for AIoT in a post-quantum world. Built and tested today: the observation and post-quantum layer "
             "(347/347 tests). Dynamic trust, quarantine and recovery are the next phase. The ESP32 uses HMAC-SHA256 (not post-quantum).")


# ------------------------------------------------------------------------------------------------ slide 2
def slide2(prs, a):
    s = new_slide(prs, 2)
    chrome(s, 2, "IDEA TITLE", a.team_name, extra="Q-SHIELD  ·  continuous device trust for AIoT")
    text(s, 0.6, 0.85, 12.1, 0.5, [("Authentication answers:  ", {"color": MUT}), ("WHO are you?", {"bold": True})], size=22, anchor="m")
    text(s, 0.6, 1.32, 12.1, 0.8, [("Q-SHIELD asks:  ", {"color": MUT, "size": 24}), ("CAN YOU STILL BE TRUSTED?", {"color": CY, "bold": True})],
         size=38, anchor="m")
    # traditional
    panel(s, 0.6, 2.4, 3.0, 3.7, alpha=55, line=LINE)
    text(s, 0.8, 2.52, 2.6, 0.25, [("TRADITIONAL IoT", {"color": DIM, "bold": True})], size=11, font=MONO, spc=60)
    trad = [("DEVICE", None), ("AUTHENTICATE", None), ("TRUST", None), ("DONE", "trust never re-evaluated")]
    ns = [node(s, 0.95, 2.95 + i * 0.78, 2.3, 0.52, t, sub, accent=LINE2 if i < 3 else DIM, fill=PANEL, tcolor=MUT, tsize=12, ssize=10)
          for i, (t, sub) in enumerate(trad)]
    for a_, b_ in zip(ns, ns[1:]):
        connect(s, a_, b_, color=DIM, w=1.25)
    # q-shield loop
    panel(s, 3.85, 2.4, 8.88, 3.7, alpha=55, line=CY, lw=1.0, glow_=(CY, 10, 14))
    text(s, 4.05, 2.52, 3.2, 0.25, [("Q-SHIELD LOOP", {"color": CY, "bold": True})], size=11, font=MONO, spc=60)
    text(s, 4.05, 2.86, 1.2, 0.22, [("SIGNALS", {"color": DIM})], size=10, font=MONO, spc=60)
    for i, lab in enumerate(("PHYSICAL", "CYBER", "VISION", "NETWORK", "CRYPTO")):
        chip(s, 5.2 + i * 1.42, 2.83, 1.28, 0.3, lab, color=BL, size=10)
    loop = [("IDENTITY", "IMPL"), ("OBSERVE", "IMPL"), ("VERIFY", "IMPL"), ("CORRELATE", "NEXT"),
            ("TRUST", "NEXT"), ("RESPOND", "DES"), ("RECOVER", "DES"), ("RE-VERIFY", "NEXT")]
    col = {"IMPL": GN, "DES": VI, "NEXT": AM}
    xs = [5.0, 7.25, 9.5, 11.75]
    pos = [(xs[i], 3.72) for i in range(4)] + [(xs[3 - i], 5.3) for i in range(4)]   # race-track: top row left->right, bottom row right->left
    ln = [node(s, x - 0.8, y - 0.34, 1.6, 0.68, t, accent=col[st], status=st, tsize=12, dashed=(st == "NEXT"))
          for (x, y), (t, st) in zip(pos, loop)]
    for k in range(8):
        connect(s, ln[k], ln[(k + 1) % 8], color=CY, w=1.75)
    text(s, 6.2, 4.32, 4.6, 0.6, [[("CONTINUOUS TRUST", {"bold": True, "color": CY, "size": 14, "spc": 60})], [("not a one-time check", {"color": MUT, "size": 11})]],
         align="c", anchor="m")
    # required sub-headings, one short line each
    for i, (h_, t_) in enumerate((("PROPOSED SOLUTION", "Trust evaluated continuously, from many signals"),
                                  ("HOW IT SOLVES THE PROBLEM", "Verified observations feed an explainable decision"),
                                  ("WHY IT'S DIFFERENT", "Integration of PQC, AI and trust, not a new algorithm"))):
        x = 0.6 + i * 4.1
        text(s, x, 6.27, 3.95, 0.22, [(h_, {"color": CY, "bold": True})], size=10, font=MONO, spc=40)
        text(s, x, 6.5, 3.95, 0.5, [(t_, {"color": TXT})], size=12)
    notes(s, "Traditional IoT authenticates once and trusts. Q-SHIELD keeps observing, verifying and correlating. Implemented today: identity, observe, verify. "
             "Correlate, trust and re-verify are the next phase; respond and recover are designed (state machine and recovery model in the decisions document).")


# ------------------------------------------------------------------------------------------------ slide 3
def slide3(prs, a):
    s = new_slide(prs, 3)
    chrome(s, 3, "TECHNICAL APPROACH", a.team_name, extra="ARCHITECTURE / FLOW DIAGRAM")
    cam = node(s, 0.9, 1.0, 2.7, 0.5, "USB CAMERA", accent=GN, tsize=12)
    # PQC profile: vision service
    panel(s, 0.6, 1.85, 3.3, 2.05, line=VI, alpha=60, lw=1.25)
    text(s, 0.78, 1.93, 2.2, 0.22, [("PQC SECURITY PROFILE", {"color": VI, "bold": True})], size=10, font=MONO, spc=30)
    tag(s, 2.72, 1.93, "IMPL")
    vis = node(s, 0.8, 2.25, 2.9, 0.75, "VISION SERVICE", "YOLO11n → observations", accent=CY, tsize=13)
    chip(s, 0.8, 3.2, 1.4, 0.36, "ML-DSA-65 ID", color=VI, size=10.5)
    chip(s, 2.3, 3.2, 1.4, 0.36, "ML-KEM-768", color=VI, size=10.5)
    text(s, 0.8, 3.62, 2.9, 0.2, [("service identity · optional session", {"color": DIM})], size=10, font=MONO)
    link(s, (2.25, 1.55), (2.25, 1.85), color=GN, w=1.5)
    # device profile: ESP32
    panel(s, 0.6, 4.3, 3.3, 1.62, line=AM, alpha=60, lw=1.25)
    text(s, 0.78, 4.38, 3.0, 0.22, [("CURRENT DEVICE AUTH · NOT PQ", {"color": AM, "bold": True})], size=10, font=MONO, spc=30)
    esp = node(s, 0.8, 4.68, 2.9, 0.55, "ESP32 ENDPOINT", accent=LINE2, tsize=13)
    chip(s, 0.8, 5.38, 1.6, 0.34, "HMAC-SHA256", color=AM, size=10.5)
    tag(s, 2.55, 5.4, "DES")
    # gateway
    gw = node(s, 5.4, 1.95, 3.0, 3.1, "", accent=CY, fill="0A2433", lw=2.0, glow_=(CY, 14, 40), rad=0.18)
    text(s, 5.4, 2.05, 3.0, 0.3, [("Q-SHIELD GATEWAY", {"bold": True, "color": TXT})], size=15, align="c", spc=30)
    for i, t in enumerate(("Verify · ML-DSA-65 · HMAC", "Freshness · replay · scope", "FastAPI · SQLite")):
        node(s, 5.6, 2.6 + i * 0.6, 2.6, 0.46, t, accent=LINE2, fill=PANEL, tsize=11, lw=0.9)
    tag(s, 5.65, 4.55, "IMPL")
    link(s, (3.9, 3.0), (5.4, 3.0), color=CY, w=2.0, glow_=True)
    dots(s, (3.9, 3.0), (5.4, 3.0), 2)
    text(s, 3.95, 2.38, 1.4, 0.55, [[("signed +", {"color": MUT})], [("optional", {"color": MUT})], [("ML-KEM session", {"color": MUT})]], size=10, font=MONO, align="c")
    polyline(s, [(3.9, 5.05), (4.65, 5.05), (4.65, 4.35), (5.4, 4.35)], color=AM, w=1.75)
    text(s, 3.95, 5.15, 1.3, 0.22, [("HMAC tag", {"color": AM})], size=10, font=MONO)
    # outputs
    outs = [("EVENTS", "reason-coded security log", "IMPL", GN, 1.55), ("TRUST", "explainable score", "NEXT", AM, 2.55),
            ("EVIDENCE", "hash-chained, signed", "DES", VI, 3.55)]
    onodes = []
    for t, sub, st, c, y in outs:
        onodes.append(node(s, 9.25, y, 3.45, 0.82, t, sub, accent=c, status=st, tsize=13, dashed=(st == "NEXT")))
        link(s, (8.4, min(max(y + 0.41, 2.2), 4.8)), (9.25, y + 0.41), color=c, w=1.4)
    resp = node(s, 9.25, 4.75, 3.45, 0.85, "RESPONSE LOOP", "quarantine · recover · re-verify", accent=AM, status="NEXT", tsize=13, dashed=True)
    bus = 12.98
    for on in onodes:
        link(s, (12.7, on.cy), (bus, on.cy), color=LINE2, w=1.25, arrow=None)
    link(s, (bus, onodes[0].cy), (bus, resp.cy), color=LINE2, w=1.25, arrow=None)
    link(s, (bus, resp.cy), (12.72, resp.cy), color=AM, w=1.5)
    # cryptographic badges
    for i, (lab, cap, c) in enumerate((("ML-KEM-768", "key establishment", VI), ("ML-DSA-65", "signatures", VI),
                                       ("AES-256-GCM", "session protection", BL), ("HMAC-SHA256", "ESP32 auth · not PQ", AM))):
        x = 0.6 + i * 3.05
        chip(s, x, 6.05, 1.5, 0.36, lab, color=c, size=10.5)
        text(s, x + 1.6, 6.05, 1.4, 0.36, [(cap, {"color": MUT})], size=10.5, anchor="m")
    for i, (h_, t_) in enumerate((("TECH STACK", "Python · FastAPI · SQLite · YOLO11n · pqcrypto"),
                                  ("METHODOLOGY", "vertical slices · observe first · trust next"),
                                  ("WHAT YOU DEMO", "signed webcam observation · tamper 401 · replay 409"))):
        x = 0.6 + i * 4.1
        text(s, x, 6.52, 3.95, 0.2, [(h_, {"color": CY, "bold": True})], size=10, font=MONO, spc=40)
        text(s, x, 6.72, 4.0, 0.3, [(t_, {"color": TXT})], size=11)
    notes(s, "Two security profiles. Vision service: ML-DSA-65 signatures, optional ML-KEM-768 session with AES-256-GCM (implemented, tested). "
             "ESP32: HMAC-SHA256 device authentication, NOT post-quantum; the gateway side is implemented and tested with a software device agent, "
             "the firmware is an unvalidated skeleton. Events are implemented; trust and the response loop are the next phase; evidence chain is designed.")


# ------------------------------------------------------------------------------------------------ slide 4
def slide4(prs, a):
    s = new_slide(prs, 4)
    chrome(s, 4, "FEASIBILITY AND VIABILITY", a.team_name, extra="validated software security path · current prototype")
    for i, (num, lab, sub) in enumerate((("347", "TESTS PASSING", "347 of 347"), ("35", "ML-KEM-768 NIST DECAP CASES", "all agree with ACVP"),
                                         ("15", "ML-DSA-65 VERIFICATION CASES", "3 valid · 12 invalid"), ("11", "REAL SIGNED OBSERVATIONS", "re-verified offline"))):
        x = 0.6 + i * 3.09
        panel(s, x, 0.85, 2.85, 1.6, alpha=60, line=LINE2, glow_=(CY, 8, 10))
        text(s, x + 0.2, 0.9, 2.5, 0.95, [(num, {"color": CY, "bold": True})], size=52, font=MONO, anchor="m")
        text(s, x + 0.2, 1.82, 2.5, 0.25, [(lab, {"color": TXT, "bold": True})], size=10.5, font=MONO)
        text(s, x + 0.2, 2.08, 2.5, 0.25, [(sub, {"color": MUT})], size=10.5)
    text(s, 0.6, 2.62, 8, 0.22, [("LIVE VALIDATION PIPELINE  ·  REAL RUN", {"color": CY, "bold": True})], size=10, font=MONO, spc=60)
    steps = [("REAL WEBCAM", "640×480"), ("YOLO11n", "observations"), ("ML-DSA SIGNED", "observation"), ("ML-KEM SESSION", "AES-256-GCM"),
             ("FASTAPI GATEWAY", "verify → store"), ("VALIDATED EVENT", "ML-DSA-65:vision-1")]
    pn = []
    for i, (t, sub) in enumerate(steps):
        pn.append(node(s, 0.6 + i * 2.07, 2.92, 1.75, 0.78, t, sub, accent=GN, tsize=11, glow_=(GN, 8, 30) if i == 5 else None))
    for a_, b_ in zip(pn, pn[1:]):
        connect(s, a_, b_, color=GN, w=1.6)
    # attack tests
    text(s, 0.6, 3.98, 6, 0.22, [("SAME PATH UNDER ATTACK  ·  CONTROLLED LOCAL TESTS", {"color": RD, "bold": True})], size=10, font=MONO, spc=40)
    rows = [("MODIFIED PAYLOAD", "signature fails", "401 REJECTED", 4.3), ("EXACT REPLAY", "id already seen", "409 REJECTED", 4.95)]
    for lab, mid, res, y in rows:
        a1 = node(s, 0.6, y, 1.85, 0.52, lab, accent=RD, tsize=11, fill="1A0F16")
        m1 = node(s, 2.85, y, 1.55, 0.52, mid, accent=LINE2, tsize=10.5, mono=True)
        r1 = node(s, 4.8, y, 1.55, 0.52, res, accent=RD, tsize=11, fill="1A0F16", glow_=(RD, 6, 25))
        connect(s, a1, m1, color=RD, w=1.4), connect(s, m1, r1, color=RD, w=1.4)
    # measured crypto cost
    panel(s, 6.75, 3.98, 5.98, 1.55, alpha=55, line=LINE2)
    text(s, 6.95, 4.04, 5.6, 0.22, [("MEASURED  ·  LAPTOP i7-13700HX  ·  NOT ESP32", {"color": DIM, "bold": True})], size=10, font=MONO)
    tiles = [("0.060", "ML-KEM keygen"), ("0.069", "encapsulate"), ("0.091", "decapsulate"), ("9.42", "ML-DSA sign"), ("0.188", "ML-DSA verify")]
    for i, (v, l) in enumerate(tiles):
        x = 6.95 + i * 1.14
        text(s, x, 4.35, 1.1, 0.5, [(v, {"color": CY if i < 3 else VI, "bold": True})], size=22, font=MONO, anchor="m")
        text(s, x, 4.86, 1.1, 0.4, [(l, {"color": MUT})], size=10)
    text(s, 6.95, 5.25, 5.6, 0.22, [("milliseconds · median of 1,500 calls per operation", {"color": DIM})], size=10, font=MONO)
    # required sub-headings
    cols = [("CHALLENGES & RISKS", AM, ["ESP32 board + firmware not validated", "ESP32-side PQC unproven", "AGPL-3.0 model · PQC library not audited"]),
            ("MITIGATION", GN, ["Software device agent unblocks work", "Swappable detector + crypto interfaces", "Trust model specified before code"]),
            ("ONE-DAY EXECUTION PLAN", CY, ["PHASE 1  Build & explore: Phases 0-3 done", "PHASE 2  Demo evaluation: signed observation", "PHASE 3  Final jury pitch: evidence + roadmap"])]
    for i, (h_, c, items) in enumerate(cols):
        x = 0.6 + i * 4.1
        text(s, x, 5.75, 3.95, 0.22, [(h_, {"color": c, "bold": True})], size=10, font=MONO, spc=40)
        text(s, x, 6.0, 4.0, 1.0, items, size=11.5, color=TXT, space_after=2)
    notes(s, "Measured, not estimated. 347/347 tests. NIST ACVP: 35 ML-KEM-768 decapsulation cases and 15 ML-DSA-65 verification cases agree. "
             "11 real signed observations, re-verified offline; modified payload 401, exact replay 409 (inside the 300 s freshness window; 401 stale after it). "
             "Latencies are laptop numbers, not ESP32. No person was in view during the recorded run.")


# ------------------------------------------------------------------------------------------------ slide 5
def slide5(prs, a):
    s = new_slide(prs, 5)
    chrome(s, 5, "IMPACT AND BENEFITS", a.team_name, extra="design intent · only the marked parts exist today")
    cx, cy, rx, ry = 5.3, 4.1, 3.3, 1.95
    stages = [("DEVICE TRUSTED", "NEXT"), ("ANOMALY DETECTED", "IMPL"), ("TRUST DECREASES", "NEXT"), ("QUARANTINE", "DES"),
              ("RECOVERY", "DES"), ("RE-AUTHENTICATE", "DES"), ("TRUST RESTORED", "NEXT")]
    col = {"IMPL": GN, "DES": VI, "NEXT": AM}
    shape(s, MSO_SHAPE.OVAL, cx - rx, cy - ry, 2 * rx, 2 * ry, fill=None, line=LINE2, lw=0.75, dashed=True)
    ns = []
    for k, (t, st) in enumerate(stages):
        x, y = ring(cx, cy, rx, ry, k, 7)
        ns.append(node(s, x - 1.03, y - 0.33, 2.06, 0.66, t, accent=col[st], status=st, tsize=11.5, dashed=(st == "NEXT"),
                       glow_=(col[st], 7, 22) if st == "IMPL" else None))
    for k in range(7):
        connect(s, ns[k], ns[(k + 1) % 7], color=CY, w=1.6)
    text(s, cx - 1.55, cy - 0.75, 3.1, 1.5, [[("From static authentication", {"color": MUT, "size": 15})], [("to continuously evaluated", {"bold": True, "size": 18})],
                                             [("device trust.", {"bold": True, "size": 22, "color": CY})]], align="c", anchor="m")
    text(s, 9.9, 1.0, 2.8, 0.22, [("WHO IT HELPS", {"color": CY, "bold": True})], size=10, font=MONO, spc=60)
    for i, d in enumerate(("SMART INFRASTRUCTURE", "INDUSTRIAL IoT", "HEALTHCARE IoT", "CONNECTED DEVICES", "CRITICAL SYSTEMS")):
        chip(s, 9.9, 1.32 + i * 0.5, 2.83, 0.4, d, color=BL, size=10.5)
    text(s, 9.9, 4.05, 2.8, 0.22, [("BENEFITS", {"color": CY, "bold": True})], size=10, font=MONO, spc=60)
    bens = [("IMPL", "Forged or replayed input rejected"), ("NEXT", "Explainable trust decisions"), ("DES", "Designed, verified recovery")]
    for i, (st, t) in enumerate(bens):
        y = 4.4 + i * 0.62
        shape(s, STATUS[st][2], 9.9, y + 0.06, 0.14, 0.14, fill=STATUS[st][1])
        text(s, 10.15, y, 2.6, 0.5, [(t, {"color": TXT})], size=12)
        text(s, 10.15, y + 0.27, 2.6, 0.2, [(STATUS[st][0], {"color": STATUS[st][1], "bold": True})], size=10, font=MONO)
    legend(s, 0.6, 6.72)
    text(s, 4.6, 6.7, 8.1, 0.28, [("Anomaly detection today = flagged observations and rejected messages. Trust, quarantine, recovery: not yet built.", {"color": DIM})],
         size=10.5)
    notes(s, "Design intent for the closed loop. Implemented today: anomaly-flagged observations and rejected/forged messages (events). Designed: quarantine, "
             "recovery and re-authentication (state machine and recovery model). Next: the trust score and its decrease/restore. No impact numbers are claimed.")


# ------------------------------------------------------------------------------------------------ slide 6
def slide6(prs, a):
    s = new_slide(prs, 6)
    chrome(s, 6, "RESEARCH AND REFERENCES", a.team_name)
    cols = [(0.6, 4.45, "SOURCES BEHIND THE IDEA", "technical basis", CY, [
        ("NIST FIPS 203 · ML-KEM-768", "csrc.nist.gov/pubs/fips/203/final"), ("NIST FIPS 204 · ML-DSA-65", "csrc.nist.gov/pubs/fips/204/final"),
        ("NIST ACVP test vectors", "github.com/usnistgov/ACVP-Server"), ("pqcrypto 1.0.0 · Apache-2.0", "pypi.org/project/pqcrypto"),
        ("YOLO11n · Ultralytics · AGPL-3.0", "github.com/ultralytics/ultralytics")]),
            (5.25, 3.7, "ETHICS & DATA HANDLING", "responsible claims", GN, [
        ("No claim of ESP32-side PQC", None), ("No fabricated benchmarks", None), ("AI observations kept apart from trust decisions", None),
        ("Limitations disclosed", None), ("AGPL model dependency disclosed", None)]),
            (9.15, 3.58, "CURRENT LIMITATIONS", "stated openly", AM, [
        ("Single-host trust domain", None), ("ESP32 hardware validation pending", None), ("ESP32 PQC feasibility pending", None),
        ("Custom PQC session: no external review", None), ("Trust / recovery stages not yet built", None)])]
    for x, w, head, cap, c, items in cols:
        panel(s, x, 0.85, w, 5.2, alpha=55, line=c if c != CY else LINE2, lw=1.0)
        shape(s, MSO_SHAPE.OVAL, x + 0.22, 1.05, 0.22, 0.22, fill=None, line=c, lw=1.5)
        text(s, x + 0.55, 1.03, w - 0.7, 0.26, [(head, {"bold": True, "color": TXT})], size=11.5, font=MONO, spc=30, anchor="m")
        text(s, x + 0.55, 1.3, w - 0.7, 0.22, [(cap, {"color": DIM})], size=10.5)
        step = 0.76 if items[0][1] else 0.86
        top = 2.3 if items[0][1] else 1.78
        if items[0][1]:                                   # technical-basis badges
            for j, lab in enumerate(("ML-KEM-768", "ML-DSA-65", "ACVP", "YOLO11n")):
                chip(s, x + 0.25 + j * 1.03, 1.68, 0.96, 0.3, lab, color=VI if j < 2 else (BL if j == 2 else AM), size=10)
        for i, (t, u) in enumerate(items):
            y = top + i * step
            shape(s, MSO_SHAPE.OVAL, x + 0.26, y + 0.09, 0.1, 0.1, fill=c)
            text(s, x + 0.55, y, w - 0.75, 0.32, [(t, {"color": TXT, "bold": bool(u)})], size=12.5)
            if u:
                text(s, x + 0.55, y + 0.32, w - 0.75, 0.24, [(u, {"color": MUT})], size=10.5, font=MONO)
    panel(s, 0.6, 6.25, 12.13, 0.7, alpha=55, line=LINE)
    for lab, val, x, w in (("TEAM NAME", a.team_name, 0.8, 2.3), ("TEAM MEMBERS", a.members, 3.2, 3.7), ("SCHOOL", a.school, 7.0, 2.7), ("TEAM LEADER CONTACT", a.contact, 9.8, 2.8)):
        text(s, x, 6.3, w, 0.2, [(lab, {"color": DIM})], size=10, font=MONO)
        text(s, x, 6.53, w, 0.3, [(val, {"color": TXT, "bold": True})], size=12)
    notes(s, "Sources: NIST FIPS 203 and 204, NIST ACVP vectors, pqcrypto 1.0.0, Ultralytics YOLO11n (AGPL-3.0). Every link was checked and returns 200. "
             "Limitations are stated openly; no ESP32-side PQC claim; no fabricated results.")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--team-name", default="[TEAM NAME]")
    ap.add_argument("--members", default="[TEAM MEMBERS]")
    ap.add_argument("--school", default="[SCHOOL]")
    ap.add_argument("--contact", default="[CONTACT]")
    ap.add_argument("--ps-number", default="[NO.]")
    ap.add_argument("--track", default="[TRACK / THEME]")
    ap.add_argument("--title", default="Q-SHIELD: Quantum-Resilient AIoT Security")
    ap.add_argument("--out", default=str(ROOT / "ppt" / "Q-SHIELD_Idea_Submission_FINAL.pptx"))
    a = ap.parse_args()
    backgrounds.build_all(BGDIR)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for f in (slide1, slide2, slide3, slide4, slide5, slide6):
        f(prs, a)
    assert len(prs.slides) == TOTAL <= 6
    prs.save(a.out)
    print(f"wrote {Path(a.out).name}: {len(prs.slides)} slides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
