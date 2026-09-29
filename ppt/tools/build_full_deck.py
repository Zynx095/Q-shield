"""Build the full technical deck (own design, dark, 16:9) and the markdown companions.

    python ppt/tools/build_full_deck.py

Outputs (all generated from ppt/tools/deck_content.py):
    ppt/Q-SHIELD_full_deck.pptx
    ppt/slide-content.md
    ppt/speaker-notes.md
This is the extended technical talk. The official-template submission deck is built separately
(build_submission_deck.py) because the template limits a submission to six slides.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deck_content import APPENDIX, FULL  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
PPT = ROOT / "ppt"

BG, PANEL, TEXT, MUTED = "0B1220", "121C31", "E9F1FA", "93A7C4"
CYAN, GOLD, GREEN, BLUE, VIOLET, GREY = "35D6F0", "F2C14E", "34D399", "60A5FA", "A78BFA", "94A3B8"
TAG_COLORS = {"IMPLEMENTED": GREEN, "MEASURED": BLUE, "DESIGNED": VIOLET, "PLANNED": GOLD, "MIXED": CYAN, "CONTEXT": GREY}


def rgb(h):
    return RGBColor.from_string(h)


def bg(slide, color=BG):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def text(slide, x, y, w, h, s, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    r.font.color.rgb = rgb(color)
    return tb


def rect(slide, x, y, w, h, fill=PANEL, line=None, dashed=False, radius=True, lw=1.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dashed:
            shp.line.dash_style = 4  # MSO_LINE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def tag_pill(slide, tag, x=10.75, y=0.42, w=2.1):
    c = TAG_COLORS[tag]
    rect(slide, x, y, w, 0.36, fill=PANEL, line=c, dashed=(tag == "PLANNED"))
    text(slide, x, y, w, 0.36, tag, 11, True, c, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, "Arial")


def title_bar(slide, title, tag):
    text(slide, 0.6, 0.42, 9.9, 0.7, title, 30 if len(title) <= 40 else 25, True, TEXT, font="Arial", anchor=MSO_ANCHOR.MIDDLE)
    tag_pill(slide, tag)


def footer(slide, n):
    text(slide, 0.6, 7.05, 6, 0.3, "Q-SHIELD  |  prototype: Phases 0-3 implemented; trust engine and recovery planned", 9, False, MUTED)
    text(slide, 12.2, 7.05, 0.6, 0.3, str(n), 9, False, MUTED, PP_ALIGN.RIGHT)


def bullet_runs(paragraph, s, size, color=TEXT):
    """Colour a leading STATUS: prefix (BUILT:, PLANNED:, ...) so implemented/designed/planned stay visible."""
    prefix, sep, rest = s.partition(": ")
    key = prefix.replace(" + ", "+")
    status = {"BUILT": GREEN, "BUILT+PLANNED": CYAN, "PLANNED": GOLD, "DESIGNED": VIOLET, "MEASURED": BLUE, "NEXT": GOLD}
    if sep and key in status:
        r = paragraph.add_run()
        r.text = prefix + ":  "
        r.font.size, r.font.bold, r.font.name = Pt(size), True, "Arial"
        r.font.color.rgb = rgb(status[key])
        s = rest
    r = paragraph.add_run()
    r.text = s
    r.font.size, r.font.name = Pt(size), "Calibri"
    r.font.color.rgb = rgb(color)


def bullets(slide, items, x, y, w, h, size=22, gap=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        bullet_runs(p, it, size)
    return tb


def callout(slide, s, y=6.05, color=GOLD):
    rect(slide, 0.6, y, 12.1, 0.75, fill=PANEL, line=color, dashed=(color == GOLD))
    text(slide, 0.85, y, 11.6, 0.75, s, 16, True, color, anchor=MSO_ANCHOR.MIDDLE)


def notes_text(sl):
    n = sl["notes"]
    lines = [f"[{sl['tag']}]  WHAT THIS SLIDE PROVES: {n['proves']}", "", "SAY:"]
    lines += [f"- {s}" for s in n["say"]]
    if n["ask"]:
        lines += ["", "LIKELY JUDGE QUESTIONS:"] + [f"Q: {q}\nA: {a}" for q, a in n["ask"]]
    if n["terms"]:
        lines += ["", f"TERMINOLOGY: {n['terms']}"]
    lines += ["", f"ACKNOWLEDGE: {n['acknowledge']}"]
    if sl.get("sources"):
        lines += ["", "SOURCES: " + "; ".join(sl["sources"])]
    return "\n".join(lines)


def add_slide(prs, sl, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    k, tag = sl["kind"], sl["tag"]
    if k == "title":
        text(s, 0.9, 2.0, 11.5, 1.2, sl["title"], 66, True, TEXT, font="Arial")
        text(s, 0.9, 3.35, 11.5, 0.8, sl["subtitle"], 28, False, CYAN)
        rect(s, 0.9, 4.55, 11.5, 0.9, fill=PANEL, line=GOLD)
        text(s, 1.15, 4.55, 11.0, 0.9, sl["line"], 17, False, TEXT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 0.9, 6.55, 11.5, 0.4, "Team: [TEAM NAME]   |   Members: [TEAM MEMBERS]   |   [SCHOOL]", 13, False, MUTED)
        tag_pill(s, tag, x=10.75, y=0.5)
    elif k == "bullets":
        title_bar(s, sl["title"], tag)
        top = 1.55
        if sl.get("headline"):
            text(s, 0.9, 1.5, 11.6, 1.3, sl["headline"], 32, True, CYAN, font="Arial")
            top = 3.05
        bullets(s, sl["bullets"], 0.9, top, 11.6, 3.0, 24 if len(sl["bullets"]) <= 4 else 21)
        if sl.get("callout"):
            callout(s, sl["callout"])
        footer(s, number)
    elif k == "columns":
        title_bar(s, sl["title"], tag)
        cols = sl["columns"]
        n = len(cols)
        gap = 0.3
        w = (12.1 - gap * (n - 1)) / n
        for i, (h, items) in enumerate(cols):
            x = 0.6 + i * (w + gap)
            rect(s, x, 1.5, w, 4.35, fill=PANEL, line=(GREEN if "Q-SHIELD" in h else MUTED))
            text(s, x + 0.25, 1.7, w - 0.5, 0.7, h, 19 if n <= 2 else 24, True, CYAN if "Q-SHIELD" in h else TEXT, font="Arial")
            bullets(s, items, x + 0.25, 2.4 if n <= 2 else 2.75, w - 0.5, 3.3, 16 if n <= 2 else 20, gap=10 if n <= 2 else 22)
        if sl.get("callout"):
            callout(s, sl["callout"])
        footer(s, number)
    elif k == "table":
        title_bar(s, sl["title"], tag)
        head, rows = sl["header"], sl["rows"]
        two = head[2] == ""
        widths = [3.0, 9.1] if two else [2.3, 4.9, 4.9]
        if two:
            head = head[:2]
        x0, y0, rh = 0.6, 1.4, 0.5 if len(rows) > 6 else 0.62
        xs = [x0]
        for w_ in widths[:-1]:
            xs.append(xs[-1] + w_)
        for j, h_ in enumerate(head):
            rect(s, xs[j], y0, widths[j] - 0.06, rh - 0.06, fill="0F2A3A", line=CYAN)
            text(s, xs[j] + 0.15, y0, widths[j] - 0.3, rh - 0.06, h_, 15, True, CYAN, anchor=MSO_ANCHOR.MIDDLE, font="Arial")
        for i, row in enumerate(rows):
            y = y0 + (i + 1) * rh
            for j in range(len(widths)):
                cell = row[j]
                warn = cell.startswith("NO") or "not " in cell.lower() and j > 0 and not two
                rect(s, xs[j], y, widths[j] - 0.06, rh - 0.06, fill=PANEL, line=None)
                col = MUTED if j == 0 else (GOLD if cell.startswith("NO") else TEXT)
                text(s, xs[j] + 0.15, y, widths[j] - 0.3, rh - 0.06, cell, 13 if len(cell) > 60 else 14, j == 0 or cell.startswith("NO"), col, anchor=MSO_ANCHOR.MIDDLE)
        if sl.get("callout"):
            callout(s, sl["callout"], y=6.15)
        footer(s, number)
    elif k == "steps":
        title_bar(s, sl["title"], tag)
        text(s, 0.6, 1.4, 12.1, 0.6, sl["scope"], 20, False, MUTED)
        for i, (num, label) in enumerate(sl["steps"]):
            col, row = i % 4, i // 4
            x, y = 0.6 + col * 3.07, 2.3 + row * 1.75
            rect(s, x, y, 2.85, 1.4, fill="2A2410", line=GOLD, dashed=True)
            text(s, x + 0.2, y + 0.12, 0.6, 0.5, num, 26, True, GOLD, font="Arial")
            text(s, x + 0.2, y + 0.65, 2.5, 0.7, label, 17, True, TEXT, font="Arial")
            if col < 3:
                text(s, x + 2.85, y + 0.45, 0.22, 0.5, ">", 18, True, GOLD, PP_ALIGN.CENTER, font="Arial")
        callout(s, sl["callout"], y=5.85)
        footer(s, number)
    elif k == "rows":
        title_bar(s, sl["title"], tag)
        items = sl["items"]
        rh = min(0.86, 5.1 / len(items))
        for i, it in enumerate(items):
            y = 1.45 + i * (rh + 0.08)
            rect(s, 0.6, y, 12.1, rh, fill=PANEL, line=None)
            rect(s, 0.6, y, 0.09, rh, fill=GOLD, line=None, radius=False)
            text(s, 0.95, y, 11.6, rh, it, 20 if len(items) <= 6 else 18, False, TEXT, anchor=MSO_ANCHOR.MIDDLE)
        footer(s, number)
    elif k == "figure":
        s.shapes.add_picture(str(ROOT / sl["figure"]), 0, 0, prs.slide_width, prs.slide_height)
        text(s, 12.2, 0.12, 0.8, 0.25, str(number), 9, False, MUTED, PP_ALIGN.RIGHT)
    elif k == "figure_text":
        title_bar(s, sl["title"], tag)
        bullets(s, sl["bullets"], 0.6, 1.5, 5.5, 5.2, 16, gap=12)
        s.shapes.add_picture(str(ROOT / sl["figure"]), Inches(6.4), Inches(1.6), width=Inches(6.5))
        footer(s, number)
    s.notes_slide.notes_text_frame.text = notes_text(sl)
    return s


def divider(prs, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    text(s, 0.9, 3.0, 11, 1.0, "Appendix", 54, True, TEXT, font="Arial")
    text(s, 0.9, 4.1, 11, 0.6, "Backup slides: authenticity vs trust, test counts, security domains, threat model, further limitations", 18, False, MUTED)
    s.notes_slide.notes_text_frame.text = "Backup section. Use only if asked."
    return s


def write_markdown(slides):
    sc = ["# Q-SHIELD: slide content (full technical deck)", "",
          "Generated by `ppt/tools/build_full_deck.py` from `ppt/tools/deck_content.py`. Do not edit by hand.",
          "Status tags: IMPLEMENTED (code + tests or recorded run) | MEASURED | DESIGNED (written only) | PLANNED (not started) | MIXED | CONTEXT.",
          "The official-template idea submission (6 slides) is a separate file: see `README.md`.", ""]
    sn = ["# Q-SHIELD: speaker notes", "", "Generated with slide-content.md. Every slide: what it proves, what to say, likely questions, terminology, what to acknowledge.", ""]
    for num, sl in slides:
        title = sl.get("title") or Path(sl.get("figure", "")).stem.replace("_", " ")
        sc += [f"## Slide {num}: {title}  [{sl['tag']}]"]
        if sl.get("figure"):
            sc.append(f"- Visual: `{sl['figure']}`")
        for b in sl.get("bullets", []):
            sc.append(f"- {b}")
        for h, items in sl.get("columns", []):
            sc.append(f"- **{h}**: " + "; ".join(items))
        if sl.get("kind") == "table":
            sc.append("| " + " | ".join(x for x in sl["header"] if x) + " |")
            sc.append("|" + "---|" * len([x for x in sl["header"] if x]))
            for r in sl["rows"]:
                sc.append("| " + " | ".join(x for x in r if x != "" or False) + " |")
        for b in sl.get("items", []):
            sc.append(f"- {b}")
        if sl.get("scope"):
            sc.append(f"- {sl['scope']}")
        for num, label in sl.get("steps", []):
            sc.append(f"  {num}. {label}")
        if sl.get("subtitle"):
            sc.append(f"- {sl['subtitle']}")
        if sl.get("line"):
            sc.append(f"- {sl['line']}")
        if sl.get("headline"):
            sc.append(f"- Headline: {sl['headline']}")
        if sl.get("callout"):
            sc.append(f"- Callout: {sl['callout']}")
        sc.append(f"- Sources: {', '.join(sl.get('sources', []))}")
        sc.append("")
        sn += [f"## Slide {num}: {title}", "", "```", notes_text(sl), "```", ""]
    (PPT / "slide-content.md").write_text("\n".join(sc), encoding="utf-8")
    (PPT / "speaker-notes.md").write_text("\n".join(sn), encoding="utf-8")


def main() -> int:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    numbered = []
    n = 0
    for sl in FULL:
        n += 1
        add_slide(prs, sl, n)
        numbered.append((n, sl))
    n += 1
    divider(prs, n)
    for sl in APPENDIX:
        n += 1
        add_slide(prs, sl, n)
        numbered.append((sl["id"], sl))
    out = PPT / "Q-SHIELD_full_deck.pptx"
    prs.save(out)
    write_markdown(numbered)
    print(f"wrote {out.name}: {len(FULL)} main slides + appendix divider + {len(APPENDIX)} backup = {len(prs.slides)} slides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
