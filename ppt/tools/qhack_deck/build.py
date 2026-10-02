"""Build the Q-SHIELD deck for Q-Hack India 2026 (Ramaiah Institute of Technology) on its official template.

    python ppt/tools/qhack_deck/build.py
    powershell -File ppt/tools/render_deck.ps1 -Pptx ppt/final/Ramaiah_Hackathon/Q-SHIELD_QHackIndia2026.pptx \
        -OutDir ppt/final/Ramaiah_Hackathon/preview -Pdf ppt/final/Ramaiah_Hackathon/Q-SHIELD_QHackIndia2026.pdf

The template (ppt/Q-Hack_India26.pptx) is opened read-only. Every slide is a clone of one of its own pages: the
background colour, the sponsor logos, the clouds, the bird illustrations, the pink section pill and the
"Q-Hack India 2026" footer are the organisers' artwork. Its ten section titles are kept word for word. The deck has
at most 12 slides (the extra two continue "Implementation and Architecture"). Text uses the template's own embedded
font, IBM Plex Sans; technical strings use Consolas. Numbers come from the repository; sources are in the notes.
"""
from __future__ import annotations

import copy
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "quantamaze_deck"))
from kit import box, clone_slide, drop_slides, elbow, line, node, notes, rule, text  # noqa: E402

ROOT = Path(__file__).resolve().parents[3]
TEMPLATE = ROOT / "ppt" / "Q-Hack_India26.pptx"
OUT_DIR = ROOT / "ppt" / "final" / "Ramaiah_Hackathon"
OUT = OUT_DIR / "Q-SHIELD_QHackIndia2026.pptx"
MAX_SLIDES = 12

PLEX = "IBM Plex Sans"
MONO = "Consolas"
INK = "111111"
MUTED = "454545"
PINK = "FA7CB5"
PINK_T = "FDE3EF"
PURPLE = "4C1D95"      # text accent (deep violet), readable on grey and lilac
DARK = "1E1E1E"
WHITE = "FFFFFF"
GREY = "E0E0E0"
LINE = "8A8A8A"
L, R = 1.3, 17.7       # content frame: clear of the clouds, the right-hand cloud and the bird illustration
B = {"bold": True}
M = {"font": MONO}

# template pages (0-based)
TITLE, PROBLEM, MATTERS, SOLUTION, QUANTUM, IMPL, PROTO, RESULTS, ROADMAP, REFS = range(10)


# ---------------------------------------------------------------------------------------------- helpers
def T(slide, x, y, w, h, paras, size=18, color=INK, **kw):
    return text(slide, x, y, w, h, paras, size=size, color=color, font=PLEX, **kw)


def BX(slide, x, y, w, h, label=None, sub=None, bold=True, sub_mono=False, **kw):
    kw.setdefault("font", PLEX)
    kw.setdefault("sub_font", MONO if sub_mono else PLEX)
    kw.setdefault("color", INK)
    kw.setdefault("sub_color", MUTED)
    kw.setdefault("line", None)
    kw.setdefault("fill", WHITE)
    kw.setdefault("radius", 0.22)
    shp = box(slide, x, y, w, h, label, sub, **kw)
    if bold and label:
        n = label.count("\n") + 1
        for p in list(shp.text_frame.paragraphs)[:n]:
            for r in p.runs:
                r.font.bold = True
    return shp


def page(prs, idx):
    """Clone template page `idx`: its background, artwork and footer; not its prompt text."""
    src = prs.slides[idx]

    def keep(shape):
        if shape.shape_type == 17:
            return shape.text_frame.text.strip() == "Q-Hack India 2026"
        return True
    s = clone_slide(prs, src, keep=keep)
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        s._element.cSld.insert(0, copy.deepcopy(bg))
    return s


def pill_of(slide):
    for sh in slide.shapes:
        if sh.shape_type == 6 and abs(sh.top / 914400 - 2.08) < 0.1 and abs(sh.height / 914400 - 0.70) < 0.05:
            return sh.left / 914400, sh.top / 914400, sh.width / 914400, sh.height / 914400
    raise RuntimeError("section pill not found")


def section(slide, title, subtitle=None):
    x, y, w, h = pill_of(slide)
    T(slide, x - 0.6, y, w + 1.2, h, [[(title, {})]], size=24, align="c", anchor="m")
    bg = slide._element.cSld.find(qn("p:bg"))
    lilac = bg is not None and b"B48AF7" in __import__("lxml.etree", fromlist=["tostring"]).tostring(bg)
    if subtitle:
        T(slide, L, 2.88, R - L, 0.5, subtitle, size=19, color=INK if lilac else MUTED, align="c")


def flow_down(slide, items, x, y, w, h, pitch, **kw):
    for i, (label, sub, st) in enumerate(items):
        BX(slide, x, y + i * pitch, w, h, label, sub, **{**kw, **st})
        if i < len(items) - 1:
            line(slide, x + w / 2, y + i * pitch + h + 0.02, x + w / 2, y + (i + 1) * pitch - 0.02, lw=1.5)


def grid_table(slide, x, y, widths, rows, head, row_h, size=16, head_size=14, highlight=()):
    cx = [x]
    for w in widths[:-1]:
        cx.append(cx[-1] + w)
    total = sum(widths)
    for hx, w, hd in zip(cx, widths, head):
        T(slide, hx, y, w - 0.1, 0.4, [[(hd, B)]], size=head_size, color=MUTED)
    rule(slide, x, y + 0.45, total, color=INK, lw=1.25)
    yy = y + 0.5
    for r, row in enumerate(rows):
        if r in highlight:
            BX(slide, x - 0.05, yy + 0.04, total + 0.1, row_h - 0.08, None, fill=PINK_T, radius=0.15)
        for i, (hx, w) in enumerate(zip(cx, widths)):
            cell = row[i]
            T(slide, hx, yy, w - 0.15, row_h, cell if isinstance(cell, list) else [[(cell, B if i == 0 else {})]],
              size=size, anchor="m", spacing=1.0, after=0)
        yy += row_h
        rule(slide, x, yy, total, color=LINE, lw=0.75)
    return yy


def bullets(items, size=16, after=4, color=INK):
    return [{"runs": it if isinstance(it, list) else [(it, {})], "bullet": True, "bullet_color": PINK, "size": size,
             "after": after, "color": color} for it in items]


# ---------------------------------------------------------------------------------------------- slides
def s01_title(prs):
    s = page(prs, TITLE)
    pills = [(0.54, 4.70, 3.38, 0.67, "Team Name:"), (0.54, 5.74, 4.53, 0.67, "Team Lead Name:"),
             (0.54, 6.73, 2.52, 0.67, "Track:"), (0.54, 7.71, 5.89, 0.70, "Problem Statement title:")]
    for x, y, w, h, lab in pills:
        T(s, x + 0.25, y, w - 0.3, h, [[(lab, {})]], size=24, anchor="m")
        if lab != "Problem Statement title:":
            T(s, x + w + 0.25, y, 5.0, h, [[("[ to fill ]", {})]], size=24, color=MUTED, anchor="m")
    T(s, 0.6, 8.55, 12.0, 1.85, [
        {"runs": [("Q-SHIELD: continuous device trust with post-quantum-signed evidence", B)], "size": 27, "after": 6},
        {"runs": [("Authentication answers ", {}), ("who are you", B), (" once. Q-SHIELD keeps asking ", {}),
                  ("can you still be trusted", {"bold": True, "color": PURPLE}), (".", {})], "size": 19},
    ], spacing=1.0)
    notes(s, "Team name, lead name and track are left as [ to fill ]: they were not supplied and must not be invented. "
             "Problem statement title: Q-SHIELD, continuous device trust with post-quantum-signed evidence.")


def s02_problem(prs):
    s = page(prs, PROBLEM)
    section(s, "Problem", "A device that authenticated can still be compromised, and nothing asks again")
    ya = 4.75
    line(s, 1.4, ya, 17.6, ya, color=INK, lw=2.0)
    node(s, 1.85, ya, r=0.15, fill=INK)
    T(s, 1.05, 3.55, 1.6, 0.95, [[("t0", M)], [("authenticate", B)]], size=17, align="c", spacing=1.0, after=0)
    T(s, 1.05, 4.92, 1.6, 0.4, "trusted", size=15, color=MUTED, align="c")
    for label, ex in (("enclosure\nopened", 5.0), ("configuration\ndrift", 8.4), ("camera covered\nor turned", 11.8),
                      ("forged or\nreplayed", 15.2)):
        node(s, ex, ya, r=0.15, fill=PINK)
        T(s, ex - 1.4, 3.6, 2.8, 0.95, [[(t, B)] for t in label.split("\n")], size=18, align="c", spacing=1.0, after=0)
        T(s, ex - 1.4, 4.92, 2.8, 0.4, [[("still trusted", M)]], size=15, color=MUTED, align="c")
    yb = 5.55
    line(s, 4.85, yb, 15.35, yb, color=PINK, lw=3.5, arrow=False)
    line(s, 4.85, yb - 0.22, 4.85, yb, color=PINK, lw=3.5, arrow=False)
    line(s, 15.35, yb - 0.22, 15.35, yb, color=PINK, lw=3.5, arrow=False)
    T(s, 4.85, 5.63, 10.5, 0.45, [[("trust should change here", {"bold": True, "color": PURPLE}),
                                   (", but authentication never asks again", {"color": MUTED})]], size=18, align="c")
    cards = [("Context", "AIoT devices at the edge (sensors, cameras, controllers) authenticate once with a provisioned "
                         "key, then run unattended."),
             ("The gap", "Authentication proves who is talking, not whether the device is still in a trustworthy "
                         "state."),
             ("The quantum angle", "Evidence signed today with RSA or ECDSA could be forged once large quantum "
                                   "computers exist.")]
    for i, (h, b) in enumerate(cards):
        x = 1.4 + i * 5.55
        BX(s, x, 6.45, 5.2, 2.5, None, fill=WHITE, radius=0.12)
        T(s, x + 0.3, 6.6, 4.65, 2.3, [[(h, {"bold": True, "color": PURPLE})], [(b, {})]], size=18, after=6)
    notes(s, "Every compromise on the timeline happens after a successful login, so the device keeps its access. "
             "The quantum angle is the motivation for post-quantum signatures on the evidence path (next slides).")


def s03_matters(prs):
    s = page(prs, MATTERS)
    section(s, "Why It Matters", "Who is exposed, and what each attack looks like with and without Q-SHIELD")
    T(s, L, 3.55, 5.0, 0.45, [[("Who faces it", B)]], size=20)
    T(s, L, 4.05, 5.0, 2.6, bullets(["Research labs and server rooms", "Smart buildings and campuses",
                                     "Critical-infrastructure sites", "Any fleet of unattended AIoT devices"], size=17),
      size=17)
    T(s, L, 6.55, 5.0, 0.45, [[("Why now", B)]], size=20)
    T(s, L, 7.05, 4.95, 2.4, [[("NIST published ", {}), ("FIPS 203", B), (" (ML-KEM) and ", {}), ("FIPS 204", B),
                               (" (ML-DSA) in August 2024: evidence signed today should already be post-quantum.", {})]],
      size=17)
    rows = [
        ("Forged or replayed message", "rejected, then forgotten", "bounded pressure (≤ 25 points); 3 replays in 10 min hold the score at 65"),
        ("Enclosure opened", "nothing to see", "physical factor CRITICAL; score held at 55"),
        ("Camera covered, frozen or turned", "nothing to see", "signed camera-health report; visual factor HIGH"),
        ("Tamper + camera or visual evidence", "nothing to see", [[("confirmed incident: score 30, ", B),
                                                                  ("quarantine", {"bold": True, "color": PURPLE})]]),
        ("Quantum adversary recording traffic", "classical signatures forgeable later",
         "ML-DSA-65 signs the evidence; ML-KEM-768 keys the session"),
    ]
    grid_table(s, 6.75, 3.55, [3.7, 2.55, 4.7], rows, ["Threat", "Authentication alone", "With Q-SHIELD"],
               row_h=0.82, size=16, highlight=(3,))
    BX(s, 6.75, 8.35, 10.95, 0.8, "No single signal quarantines a device. Two independent authenticated signals within 60 s do.",
       fill=PINK, radius=0.5, size=17)
    notes(s, "Threat model from docs/architecture/trust-engine.md (sections 4.5, 5, 9). FIPS 203 and FIPS 204 were "
             "published by NIST in August 2024.")


def s04_solution(prs):
    s = page(prs, SOLUTION)
    section(s, "Proposed Solution", "A closed loop around every device: trust is computed continuously, enforced, and earned back")
    cx, cy, rx, ry = 6.2, 6.55, 3.5, 2.1
    steps = [("OBSERVE", "device and camera"), ("VERIFY", "HMAC, ML-DSA-65"), ("SCORE", "trust engine"),
             ("ENFORCE", "quarantine at gateway"), ("RECOVER", "verified, earned"), ("PROVE", "signed evidence chain")]
    pos = [(cx + rx * math.cos(math.radians(90 - i * 60)), cy - ry * math.sin(math.radians(90 - i * 60))) for i in range(6)]

    def edge(px, py, dx, dy, hw=1.35, hh=0.45):
        t = min(hw / abs(dx) if dx else 1e9, hh / abs(dy) if dy else 1e9)
        return px + dx * t, py + dy * t
    for i in range(6):
        (px, py), (nx, ny) = pos[i], pos[(i + 1) % 6]
        d = math.hypot(nx - px, ny - py)
        ux, uy = (nx - px) / d, (ny - py) / d
        ax, ay = edge(px, py, ux, uy)
        bx, by = edge(nx, ny, -ux, -uy)
        line(s, ax + ux * 0.06, ay + uy * 0.06, bx - ux * 0.08, by - uy * 0.08, color=INK, lw=1.5)
    for (px, py), (lab, sub) in zip(pos, steps):
        hot = lab in ("ENFORCE", "RECOVER")
        BX(s, px - 1.35, py - 0.45, 2.7, 0.9, lab, sub, fill=PINK if hot else WHITE, size=18, sub_size=14,
           sub_color=INK if hot else MUTED, radius=0.5)
    T(s, cx - 1.1, cy - 0.45, 2.2, 0.9, [[("continuous", B)], [("trust", B)]], size=19, color=PURPLE, align="c",
      anchor="m", spacing=1.0, after=0)
    T(s, 11.0, 3.6, 6.7, 0.45, [[("What we are building", B)]], size=20)
    items = [("Continuous six-factor trust engine", "every score change lists the reasons that sum to it"),
             ("The camera as part of the security boundary", "covered, frozen or turned is signed evidence"),
             ("Quarantine enforced at the gateway", "the device's normal channel is refused (403)"),
             ("Verified, earned recovery", "health checks first, then trust rebuilt from fresh evidence"),
             ("Post-quantum evidence", "ML-DSA-65 signatures, ML-KEM-768 sessions")]
    for i, (a, b) in enumerate(items):
        y = 4.2 + i * 1.0
        node(s, 11.15, y + 0.2, r=0.11, fill=PINK, line=INK, lw=1.0)
        T(s, 11.45, y, 6.25, 0.95, [[(a, B)], [(b, {"color": MUTED, "size": 15})]], size=17, spacing=1.0, after=0)
    notes(s, "The novelty is the loop, not one algorithm. Pink marks the two steps one-time authentication lacks: "
             "enforcement and earned recovery.")


def s05_quantum(prs):
    s = page(prs, QUANTUM)
    section(s, "Why Quantum?", "Quantum computing is the threat model; post-quantum cryptography is how Q-SHIELD answers it")
    heads = [("Quantum capability", 1.4, 3.4), ("What it breaks", 5.25, 4.6), ("Q-SHIELD's answer", 10.3, 7.3)]
    for h, x, w in heads:
        T(s, x, 3.5, w, 0.4, [[(h, B)]], size=15, color=MUTED)
    rows = [("Shor's algorithm", "RSA, ECDSA, ECDH:\ntoday's signatures and key exchange",
             "ML-DSA-65 (FIPS 204) signs observations and evidence;\nML-KEM-768 (FIPS 203) agrees session keys"),
            ("Grover's algorithm", "at most a square-root speed-up\non brute-force key search",
             "AES-256-GCM, SHA-256 and 256-bit HMAC keys\nkeep a large security margin"),
            ("Harvest now, forge later", "long-lived evidence\nrecorded today",
             "every evidence-chain entry is signed with ML-DSA-65")]
    for i, (a, b, c) in enumerate(rows):
        y = 3.95 + i * 1.15
        BX(s, 1.4, y, 3.4, 0.95, a, fill=DARK, color=WHITE, size=17, radius=0.3)
        line(s, 4.85, y + 0.47, 5.2, y + 0.47, lw=1.75)
        BX(s, 5.25, y, 4.6, 0.95, None, fill=WHITE, radius=0.3)
        T(s, 5.4, y, 4.35, 0.95, b, size=15, anchor="m", spacing=1.0, after=0)
        line(s, 9.9, y + 0.47, 10.25, y + 0.47, lw=1.75)
        BX(s, 10.3, y, 7.3, 0.95, None, fill=PINK_T, line=PINK, lw=1.5, radius=0.3)
        T(s, 10.5, y, 7.0, 0.95, c, size=15, anchor="m", spacing=1.0, after=0)
    T(s, 1.4, 7.55, 8.0, 0.45, [[("The cost of going post-quantum (measured, laptop CPU)", B)]], size=17)
    cost = [("signature size", "ML-DSA-65 3309 B vs ECDSA P-256 64 B"), ("sign / verify", "9.42 ms / 0.188 ms (median)"),
            ("ML-KEM-768 encapsulate", "0.069 ms (median)"), ("full session handshake", "10.0 ms (median)")]
    for i, (k, v) in enumerate(cost):
        T(s, 1.4, 8.0 + i * 0.42, 3.3, 0.42, k, size=15, color=MUTED, anchor="m")
        T(s, 4.7, 8.0 + i * 0.42, 5.2, 0.42, [[(v, M)]], size=15, anchor="m")
    BX(s, 10.3, 7.6, 7.3, 1.3, None, fill=WHITE, radius=0.15)
    T(s, 10.5, 7.68, 6.95, 1.5, [[("Honest scope: ", {"bold": True, "color": PURPLE}),
                                  ("Q-SHIELD runs no quantum circuits and no quantum hardware. Its quantum-relevant "
                                   "part is post-quantum cryptography on the evidence path; the ESP32 device path is "
                                   "HMAC-SHA256, not a post-quantum signature.", {})]], size=15, spacing=1.0)
    notes(s, "Shor's algorithm (1994) breaks factoring and discrete-log based schemes (RSA, ECDSA, ECDH). Grover's "
             "algorithm (1996) gives at most a quadratic speed-up for search, so 256-bit symmetric keys and SHA-256 "
             "keep a large margin. Numbers: docs/results/pqc-benchmark.json (i7-13700HX, pqcrypto 1.0.0, Python "
             "bindings). ECDSA P-256 raw signature: 64 bytes.")


def s06_arch(prs):
    s = page(prs, IMPL)
    section(s, "Implementation and Architecture", "1 of 3: how the parts work together, quantum-safe and classical")
    dx, dy = 0.4, -0.3
    h = 0.95

    def bx(x, y, w, hh, lab, sub=None, **kw):
        return BX(s, x + dx, y + dy, w, hh, lab, sub, line=kw.pop("line", INK), lw=kw.pop("lw", 1.25), **kw)

    def ln(x1, y1, x2, y2, **kw):
        return line(s, x1 + dx, y1 + dy, x2 + dx, y2 + dy, lw=1.75, **kw)
    bx(0.9, 3.85, 2.75, h, "ESP32 / device agent", "simulated today", size=17, sub_size=14)
    bx(4.15, 3.85, 3.1, h, "Identity", "HMAC-SHA256 + counter", size=17, sub_size=14, sub_mono=True)
    ln(3.69, 4.32, 4.11, 4.32)
    bx(0.9, 5.35, 2.75, h, "Vision service", "camera + YOLO11n", size=17, sub_size=14)
    bx(4.15, 5.35, 3.1, h, "PQC session", "ML-KEM-768, AES-256-GCM", size=17, sub_size=14, sub_mono=True,
       fill=PINK_T, line=PINK, lw=2.0)
    bx(7.75, 5.35, 3.0, h, "Signed observations", "ML-DSA-65", size=17, sub_size=14, sub_mono=True, fill=PINK_T,
       line=PINK, lw=2.0)
    ln(3.69, 5.82, 4.11, 5.82)
    ln(7.29, 5.82, 7.71, 5.82)
    bx(11.3, 3.85, 2.55, 2.45, "Gateway", "authenticate\nenforce\nrecord", size=19, sub_size=14)
    ln(7.29, 4.32, 11.26, 4.32)
    ln(10.79, 5.82, 11.26, 5.82)
    bx(14.35, 3.85, 2.75, 2.45, "Trust engine", "six factors\ncaps, decay\nstate machine", size=19, sub_size=14,
       fill=DARK, line=DARK, color=WHITE, sub_color=WHITE)
    ln(13.89, 5.07, 14.31, 5.07)
    bx(14.62, 6.68, 2.2, 1.2, "Decision", kind="diamond", size=15, inset=0.02)
    ln(15.72, 6.34, 15.72, 6.64)
    bx(14.2, 8.6, 2.9, 0.9, "TRUSTED or SUSPICIOUS", "normal access", size=15, sub_size=13)
    ln(15.72, 7.9, 15.72, 8.56)
    T(s, 15.8 + dx, 7.98 + dy, 1.3, 0.4, [[("≥ 50", M)]], size=14, color=MUTED)
    bx(11.3, 6.82, 2.55, 0.9, "QUARANTINE", "normal channel 403", fill=PINK, line=PINK, size=17, sub_size=13,
       sub_color=INK, sub_mono=True)
    ln(14.6, 7.27, 13.89, 7.27)
    T(s, 13.95 + dx, 6.82 + dy, 0.95, 0.4, [[("< 50", M)]], size=14, color=MUTED)
    bx(7.75, 6.82, 3.0, 0.9, "Recovery", "operator-started, verified", size=17, sub_size=13)
    ln(11.26, 7.27, 10.79, 7.27)
    bx(4.15, 6.82, 3.1, 0.9, "VERIFIED → TRUSTED", "trust re-earned", size=17, sub_size=13)
    ln(7.71, 7.27, 7.29, 7.27)
    bx(0.9, 8.35, 5.0, 1.0, "Digital twin", "expected vs self-reported state", line=None, size=16, sub_size=14,
       align="l", inset=0.2)
    bx(6.15, 8.35, 6.45, 1.0, "Evidence chain", "every decision SHA-256 linked, ML-DSA-65 signed", line=PINK, lw=2.0,
       size=16, sub_size=14, align="l", inset=0.2)
    comps = [("Quantum-safe", "ML-KEM-768, ML-DSA-65 (pqcrypto)", PINK_T, PINK),
             ("Classical", "FastAPI, trust engine, SQLite, YOLO11n, OpenCV, dashboard", WHITE, None),
             ("Device path", "HMAC-SHA256 with a counter (not post-quantum)", WHITE, None)]
    for i, (a, b, fill, ln_) in enumerate(comps):
        x = L + i * 5.55
        BX(s, x, 9.3, 5.3, 0.8, None, fill=fill, line=ln_, lw=1.5, radius=0.5)
        T(s, x + 0.25, 9.3, 4.95, 0.8, [[(a + ": ", B), (b, {})]], size=14, anchor="m", spacing=1.0)
    notes(s, "Pink outline = post-quantum components. The device lane authenticates with HMAC-SHA256; the vision lane "
             "uses an ML-KEM-768 session carrying ML-DSA-65 signed observations. The trust engine decides; the gateway "
             "enforces. Sources: backend/api/app.py, backend/trust/, backend/security/session.py.")


def s07_pqc(prs):
    s = page(prs, IMPL)
    section(s, "Implementation and Architecture", "2 of 3: the post-quantum evidence path")
    T(s, L, 3.5, 7.9, 0.45, [[("ML-KEM-768 session (vision service ↔ gateway)", B)]], size=18)
    steps = [("Gateway key pair, provisioned once", "client pins the public key, 1184 B"),
             ("Client encapsulates", "→ shared secret + ciphertext, 1088 B"),
             ("Init signed with ML-DSA-65; gateway verifies, decapsulates", "→ shared secret, 32 B"),
             ("HKDF-SHA256 on both sides", "salt = both nonces, info = transcript"),
             ("Key confirmation", "a wrong pinned key: nothing is sent"),
             ("AES-256-GCM session", "per-direction keys, 64-bit counters, AAD")]
    for i, (a, b) in enumerate(steps):
        y = 4.0 + i * 0.8
        hot = i in (2, 5)
        BX(s, L, y, 8.0, 0.68, None, fill=PINK_T if hot else WHITE, line=PINK if hot else None, lw=1.5, radius=0.3)
        BX(s, L + 0.12, y + 0.12, 0.44, 0.44, str(i + 1), fill=PINK if not hot else DARK, color=INK if not hot else WHITE,
           size=14, kind="oval")
        T(s, L + 0.7, y, 7.2, 0.68, [[(a, B)], [(b, M)]], size=14, anchor="m", spacing=1.0, after=0)
    T(s, L, 8.85, 8.0, 0.9, [[("median, laptop: ", {"color": MUTED}), ("encaps 0.069 ms, decaps 0.091 ms, handshake 10.0 ms", M)],
                             [("Composed by us from standard primitives: not TLS, not externally reviewed.", {"color": MUTED})]],
      size=14, after=2)
    x2, w2 = 9.9, 7.8
    T(s, x2, 3.5, w2, 0.45, [[("ML-DSA-65 signed observation", B)]], size=18)
    flow_down(s, [("Observation", "device, object, zone, confidence, time", {}),
                  ("Canonical encoding", "length-prefixed fields, never JSON", {}),
                  ("ML-DSA-65 sign", "context qshield/signed-observation/v1", {"fill": PINK}),
                  ("Signature, 3309 B", "bound to signer, source and allowed devices", {}),
                  ("Gateway verify", "active signer, in scope, ±300 s, not seen before", {"fill": DARK, "color": WHITE,
                                                                                          "sub_color": WHITE})],
              x2, 4.0, w2, 0.62, 0.78, size=15, sub_size=13, sub_mono=True)
    line(s, x2 + w2 / 2, 7.74, x2 + w2 / 2, 7.88, lw=1.5, arrow=False)
    line(s, x2 + 1.85, 7.88, x2 + w2 - 1.85, 7.88, lw=1.5, arrow=False)
    line(s, x2 + 1.85, 7.88, x2 + 1.85, 8.05, lw=1.5)
    line(s, x2 + w2 - 1.85, 7.88, x2 + w2 - 1.85, 8.05, lw=1.5)
    BX(s, x2, 8.08, 3.7, 0.72, "ACCEPT 200", "→ trust engine", size=15, sub_size=13, line=INK, lw=1.5)
    BX(s, x2 + 4.1, 8.08, 3.7, 0.72, "REJECT", "401 forged or stale, 409 replay", size=15, sub_size=13, fill=PINK_T,
       line=PINK, lw=1.5, sub_mono=True)
    T(s, x2, 8.95, w2, 0.85, [[("Evidence chain: ", B), ("every decision SHA-256 linked and ML-DSA-65 signed; ", {}),
                              ("26 entries VERIFIED", B), (" in the live demo.", {})]], size=14, spacing=1.0)
    notes(s, "Handshake: backend/security/session.py. Signed observations: backend/protocol/signed_observation.py. "
             "Evidence chain: backend/evidence/chain.py (context qshield/evidence/v1); not anchored externally, so "
             "cutting the tail needs an outside copy of the head. Numbers: docs/results/pqc-benchmark.json.")


def s08_trust(prs):
    s = page(prs, IMPL)
    section(s, "Implementation and Architecture", "3 of 3: the trust engine and the camera as part of the security boundary")
    T(s, L, 3.5, 7.7, 0.45, [[("Six-factor trust", B)]], size=18)
    for i, (name, w) in enumerate((("Identity / crypto", 25), ("Physical", 20), ("Configuration integrity", 15),
                                    ("Sensor consistency", 15), ("Visual", 15), ("Network liveness", 10))):
        y = 3.98 + i * 0.52
        T(s, L, y, 3.1, 0.48, name, size=15, anchor="m")
        bw = 3.4 * w / 25
        BX(s, L + 3.15, y + 0.12, bw, 0.25, None, fill=PINK, kind="rect")
        T(s, L + 3.25 + bw, y, 0.9, 0.48, [[(f"{w}%", B)]], size=15, anchor="m")
    T(s, L, 7.15, 7.8, 0.45, [[("score = min(100 − Σ wᵢ·pᵢ − pressure, caps)", M)]], size=15)
    sc = 7.7 / 100
    BX(s, L, 7.7, 50 * sc, 0.55, "QUARANTINED  < 50", fill=DARK, color=WHITE, kind="rect", size=13)
    BX(s, L + 50 * sc, 7.7, 30 * sc, 0.55, "SUSPICIOUS", fill=PINK, kind="rect", size=13)
    BX(s, L + 80 * sc, 7.7, 20 * sc, 0.55, "TRUSTED", fill=WHITE, line=INK, lw=1.5, kind="rect", size=13)
    T(s, L, 8.35, 7.8, 1.2, [[("back to TRUSTED only at 85; pressure ≤ 25, so forged traffic alone never quarantines; "
                               "only clean authenticated evidence earns recovery time", {"color": MUTED})]], size=14,
      spacing=1.0)
    x2 = 9.6
    T(s, x2, 3.5, 8.1, 0.45, [[("Camera checks (signed by the vision service)", B)]], size=18)
    rows = [("Occluded", "dark, flat or blinded for 2 s", "HIGH"), ("Moved", "shifted, or similarity < 0.7", "HIGH"),
            ("Frozen", "the same frame repeated", "HIGH"), ("Degraded", "low light or blur", "MEDIUM"),
            ("Too close", "box > 35% of frame, not a distance", "LOW"), ("Source lost", "no frames", "MEDIUM")]
    grid_table(s, x2, 3.95, [2.0, 4.6, 1.5], rows, ["State", "How it is measured", "Effect"], row_h=0.5, size=14,
               head_size=13, highlight=(0, 1, 2))
    y = 7.75
    BX(s, x2, y, 2.1, 0.75, "Tamper", "HMAC report", size=14, sub_size=12, line=INK, lw=1.25)
    T(s, x2 + 2.1, y, 0.4, 0.75, [[("+", B)]], size=20, align="c", anchor="m")
    BX(s, x2 + 2.5, y, 2.6, 0.75, "Camera interference", "signed, within 60 s", size=14, sub_size=12, fill=PINK_T,
       line=PINK, lw=1.5)
    line(s, x2 + 5.15, y + 0.37, x2 + 5.55, y + 0.37, lw=1.75)
    BX(s, x2 + 5.6, y, 2.5, 0.75, "QUARANTINE", "confirmed, score 30", size=14, sub_size=12, fill=PINK)
    T(s, x2, 8.65, 8.1, 0.9, [[("Camera alone: a penalty, never quarantine. A camera report and a detection from the "
                                "same signer count once.", {"color": MUTED})]], size=14, spacing=1.0)
    notes(s, "Weights and caps: docs/architecture/trust-engine.md (design choices, not calibrated). Camera checks: "
             "ai/vision/health.py and ai/vision/proximity.py; tested on synthetic frames; on the live USB webcam a "
             "moving person produced no false camera-health report in 2 x 90 s. Covering, turning and freezing the "
             "real camera have not been staged. Proximity is image-space: no distance is measured.")


def s09_proto(prs):
    s = page(prs, PROTO)
    section(s, "Minimal Prototype", "What we have actually built: running software, tested end to end")
    tiles = [("Gateway", "FastAPI + SQLite", "real"), ("Trust engine", "six factors, caps, states", "real"),
             ("PQC layer", "ML-KEM-768, ML-DSA-65", "real"), ("Vision service", "YOLO11n + camera checks", "real"),
             ("Quarantine", "enforced after authentication", "real"), ("Recovery", "orchestrator + deadlines", "real"),
             ("Digital twin", "expected vs reported", "real"), ("Evidence chain", "linked, signed, verified", "real"),
             ("Dashboard", "live view + presentation mode", "real"), ("Device agent", "software, labelled", "sim"),
             ("Attack simulator", "synthetic, labelled", "sim"), ("ESP32 firmware", "skeleton, not compiled", "next")]
    for i, (a, b, kind) in enumerate(tiles):
        x = L + (i % 4) * 4.15
        y = 3.5 + (i // 4) * 1.08
        fill = {"real": WHITE, "sim": PINK_T, "next": GREY}[kind]
        BX(s, x, y, 3.95, 0.95, None, fill=fill, line=None if kind == "real" else (PINK if kind == "sim" else LINE),
           lw=1.5, dash=kind == "next", radius=0.25)
        tag = {"real": "REAL", "sim": "SIMULATED", "next": "NOT YET"}[kind]
        T(s, x + 0.2, y + 0.05, 3.6, 0.85, [[(a, B), ("   " + tag, {"size": 13, "color": PURPLE, "bold": True})],
                                             [(b, {"color": MUTED, "size": 14})]], size=16, spacing=1.0, after=0)
    T(s, L, 6.85, 8.0, 0.4, [[("Self-healing recovery", B)]], size=17)
    steps = ["Quarantine", "Authorize", "Remediate", "Health checks", "Twin match", "Verified", "Recovered", "Trusted"]
    caps_ = ["", "operator", "known-good config", "3 clean reports", "expected = reported", "ramp begins", "score ≥ 50",
             "score ≥ 85"]
    for i, (a, c) in enumerate(zip(steps, caps_)):
        x = L + i * 2.07
        fill = DARK if i == 0 else (PINK if a in ("Verified", "Recovered") else (DARK if a == "Trusted" else WHITE))
        BX(s, x, 7.3, 1.9, 0.7, a, fill=fill, color=WHITE if fill == DARK else INK, size=14, radius=0.5)
        if c:
            T(s, x - 0.05, 8.03, 2.0, 0.4, c, size=12, color=MUTED, align="c")
        if i < 7:
            line(s, x + 1.92, 7.65, x + 2.05, 7.65, lw=1.25)
    T(s, L, 8.45, 16.0, 0.4, [[("any failure → QUARANTINED: ", {"bold": True, "color": PURPLE}),
                               ("an authenticated fault, the 15-min verification deadline, the 2-h ramp deadline, "
                                "revocation", {"color": MUTED})]], size=14)
    T(s, L, 9.0, 16.0, 0.5, [[("Run it:  ", B), ("python scripts/demo_full.py --webcam --pace 2 --hold", M)]], size=16)
    notes(s, "Everything marked REAL runs and is tested (785 Python and 40 JS tests). Device telemetry comes from a "
             "labelled software agent; attacks come from attack_simulation/ and are labelled; the ESP32 firmware has "
             "never been compiled. Repository: github.com/Zynx095/Q-shield.")


def s10_results(prs):
    s = page(prs, RESULTS)
    section(s, "Results / Validation", "Validated on the real gateway with a real USB webcam; attacks are simulated and labelled")
    BX(s, L, 3.5, 8.1, 4.45, None, fill=WHITE, radius=0.08)
    cx, cy, cw, ch = L + 0.75, 3.85, 7.0, 3.05

    def ys(v):
        return cy + ch - v / 100 * ch
    BX(s, cx, ys(50), cw, ch * 0.5, None, fill=GREY, kind="rect")
    BX(s, cx, ys(80), cw, ch * 0.3, None, fill=PINK_T, kind="rect")
    for v in (0, 50, 80):
        T(s, cx - 0.6, ys(v) - 0.18, 0.5, 0.36, [[(str(v), M)]], size=12, color=MUTED, align="r")
    T(s, cx + 0.08, ys(50) + 0.03, 2.5, 0.35, "quarantine", size=12, color=MUTED)
    T(s, cx + 0.08, ys(80) + 0.03, 2.5, 0.35, "suspicious", size=12, color=PURPLE)
    pts = [("join", 100), ("webcam", 90), ("forged", 80), ("replay", 72), ("tamper", 30), ("verified", 30),
           ("recovered", 80), ("trusted", 85)]
    stp = cw / (len(pts) - 1)
    co = [(cx + i * stp, ys(v)) for i, (_, v) in enumerate(pts)]
    for (x1, y1), (x2, y2) in zip(co, co[1:]):
        line(s, x1, y1, x2, y2, color=INK, lw=2.25, arrow=False)
    for (x, y_), (lab, v) in zip(co, pts):
        node(s, x, y_, r=0.1, fill=PINK if v < 50 else WHITE, line=INK, lw=1.5)
        T(s, x - 0.5, y_ - 0.46 if v > 40 else y_ + 0.1, 1.0, 0.35, [[(str(v), M)]], size=13, align="c")
        T(s, x - 0.6, cy + ch + 0.05, 1.2, 0.35, lab, size=12, color=MUTED, align="c")
    T(s, L, 7.5, 8.1, 0.4, "Live run, 2026-10-02; recovery shown in announced time-lapse", size=12, color=MUTED, align="c")
    x2 = 9.95
    outcomes = [("Forged observation", "REJECTED 401"), ("Replayed observation", "REJECTED 409"),
                ("Tamper + signed visual evidence", "QUARANTINED, score 30"), ("Forged 'all clear' in quarantine", "401"),
                ("Verified recovery", "TRUSTED 85; chain VERIFIED")]
    for i, (a, b) in enumerate(outcomes):
        y = 3.5 + i * 0.9
        BX(s, x2, y, 7.75, 0.78, None, fill=WHITE, radius=0.5)
        T(s, x2 + 0.3, y, 4.0, 0.78, [[(a, B)]], size=15, anchor="m")
        BX(s, x2 + 4.25, y + 0.12, 3.35, 0.54, b, fill=PINK if "QUARANT" in b else GREY, size=13, radius=0.5,
           font=MONO)
    tiles = [("785", "Python tests passing"), ("40", "JavaScript tests passing"), ("70", "NIST ACVP vectors agree"),
             ("10.0 ms", "full PQC handshake (median)"), ("~5 fps", "live USB webcam, no false camera alarm in 2 × 90 s")]
    for i, (big, lab) in enumerate(tiles):
        x = L + i * 3.32
        BX(s, x, 8.2, 3.12, 1.45, None, fill=WHITE, radius=0.15)
        T(s, x + 0.2, 8.25, 2.8, 0.65, [[(big, B)]], size=26)
        T(s, x + 0.2, 8.88, 2.8, 0.75, lab, size=13, color=MUTED, spacing=1.0)
    notes(s, "Measured run of scripts/demo_full.py --webcam --pace 2 --hold: 100 -> 90 (real detections) -> 80 (forged) "
             "-> 72 SUSPICIOUS (replay) -> 30 QUARANTINED (tamper + signed visual) -> VERIFIED 30 -> RECOVERED 80 -> "
             "TRUSTED 85; evidence chain 26 entries, VERIFIED. Tests: python -m pytest (785 passed), node --test (40). "
             "ACVP: pqc/README.md. Detection accuracy has not been measured.")


def s11_roadmap(prs):
    s = page(prs, ROADMAP)
    section(s, "Roadmap", "If selected for Round 2: from software prototype to a validated hardware deployment")
    stages = [("Hardware bring-up", ["compile and flash the ESP32", "match the 7 byte-exact protocol test vectors",
                                     "reed switch, then BME280 and MPU6050"]),
              ("Staged validation", ["cover, turn, freeze and unplug the camera", "open the enclosure",
                                     "record false alarms and misses"]),
              ("Calibrate", ["set camera thresholds from recorded data", "tune trust weights from incidents, not the demo"]),
              ("Harden", ["anchor the evidence head externally", "key rotation, TLS by default",
                          "more than one gateway"]),
              ("Research", ["measure ML-KEM / ML-DSA on an ESP32-class device", "a post-quantum device path if feasible"])]
    w, gap = 3.08, 0.25
    for i, (title, items) in enumerate(stages):
        x = L + i * (w + gap)
        BX(s, x, 3.6, 0.6, 0.6, str(i + 1), kind="oval", fill=PINK if i == 0 else WHITE, line=INK, lw=1.25, size=17)
        if i < 4:
            line(s, x + 0.64, 3.9, x + w + gap - 0.04, 3.9, color=INK, lw=1.5)
        T(s, x, 4.35, w, 0.5, [[(title, B)]], size=18)
        BX(s, x, 4.95, w, 2.45, None, fill=WHITE, radius=0.1)
        T(s, x + 0.15, 5.05, w - 0.25, 2.3, bullets(items, size=15, after=8), size=15)
    T(s, L, 7.75, 15.8, 0.55, [[("Hardware is an adapter step: ", {"bold": True, "color": PURPLE}),
                                ("same protocol, same gateway, same trust engine (docs/hardware/).", {})]], size=17)
    notes(s, "Plan from docs/hardware/hardware-architecture.md section 2 and docs/IMPLEMENTATION_STATUS.md 'What is "
             "left'. Nothing here has been done yet.")


def s12_refs(prs):
    s = page(prs, REFS)
    section(s, "References and Sources", "Standards, research, libraries and our repository")
    left = [("Standards", [[("NIST FIPS 203, ML-KEM (2024)  ", {}), ("doi.org/10.6028/NIST.FIPS.203", M)],
                           [("NIST FIPS 204, ML-DSA (2024)  ", {}), ("doi.org/10.6028/NIST.FIPS.204", M)],
                           [("NIST SP 800-38D, GCM  ", {}), ("doi.org/10.6028/NIST.SP.800-38D", M)],
                           [("NIST FIPS 180-4, SHA-256  ", {}), ("doi.org/10.6028/NIST.FIPS.180-4", M)],
                           [("RFC 5869 HKDF, RFC 2104 HMAC  ", {}), ("rfc-editor.org", M)],
                           [("NIST ACVP test vectors  ", {}), ("github.com/usnistgov/ACVP-Server", M)]]),
            ("Research", [[("P. W. Shor, Algorithms for quantum computation: discrete logarithms and factoring, FOCS 1994", {})],
                          [("L. K. Grover, A fast quantum mechanical algorithm for database search, STOC 1996", {})]])]
    right = [("Libraries and tools", [[("pqcrypto 1.0.0 (Apache-2.0): ML-KEM-768, ML-DSA-65", {})],
                                      [("cryptography (PyCA): AES-256-GCM, HKDF", {})],
                                      [("kyber-py, dilithium-py: test-only cross-checks", {})],
                                      [("Ultralytics YOLO11n (AGPL-3.0), OpenCV, NumPy", {})],
                                      [("FastAPI, Pydantic, SQLite", {})]]),
             ("Our work", [[("github.com/Zynx095/Q-shield", {"font": MONO, "bold": True})],
                           [("an original repository, not a fork; third-party libraries are dependencies under their "
                             "licences", {})],
                           [("docs/: specifications, test results and limitations", {})]])]
    lines = {"Standards": 6, "Research": 3, "Libraries and tools": 5, "Our work": 4}
    for col, x in ((left, L), (right, 9.75)):
        y = 3.5
        for head, items in col:
            T(s, x, y, 7.9, 0.45, [[(head, B)]], size=18)
            y += 0.5
            hgt = 0.37 * lines[head] + 0.3
            BX(s, x, y, 7.9, hgt, None, fill=WHITE, radius=0.1)
            T(s, x + 0.2, y + 0.12, 7.55, hgt, bullets(items, size=14, after=3), size=14, spacing=1.0)
            y += hgt + 0.35
    notes(s, "All links are public. The repository contains the specifications (docs/architecture), the test and "
             "benchmark results (docs/results, docs/testing.md) and the limitations (README.md).")


SLIDES = [s01_title, s02_problem, s03_matters, s04_solution, s05_quantum, s06_arch, s07_pqc, s08_trust, s09_proto,
          s10_results, s11_roadmap, s12_refs]


def main() -> int:
    assert len(SLIDES) <= MAX_SLIDES, "at most 12 slides"
    prs = Presentation(str(TEMPLATE))
    n_template = len(prs.slides)
    for fn in SLIDES:
        fn(prs)
    drop_slides(prs, n_template)
    assert len(prs.slides) == len(SLIDES) <= MAX_SLIDES
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({len(prs.slides)} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
